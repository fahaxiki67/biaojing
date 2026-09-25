# -*- coding: utf-8 -*-
"""标镜多行清单报价的证据注册表回归（端到端发现的真实缺陷）。

口径：确认 price_lines 时，facts 证据注册表必须包含**每一行**的单价
证据与字段证据 ID——装配只登记确认行顶层 evidence_id 会让第 2 行起的
清单行在规则端全部"证据缺失/无效"被排除，多行清单（真实标书常态）
发生静默覆盖损失。

运行：
    cd app && python3 -W error::ResourceWarning -m unittest \\
        biaojing.tests.test_price_lines_evidence_registry -v
"""

from __future__ import annotations

import http.client
import io
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.parse
import urllib.request

import openpyxl

from biaojing import webapp, workspace


class PriceLinesEvidenceRegistryTests(unittest.TestCase):

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="biaojing_pl_evreg_")
        self.wb = workspace.Workbench(self.temp.name)
        self.srv = webapp.WorkbenchServer(("127.0.0.1", 0), self.wb)
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}"
        self.thread = threading.Thread(target=self.srv.serve_forever,
                                       daemon=True)
        self.thread.start()

    def tearDown(self):
        self.srv.shutdown()
        self.srv.server_close()
        self.wb.close()
        self.temp.cleanup()

    def _post_json(self, path: str, payload: dict) -> dict:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(
            self.base + path, data=body, method="POST",
            headers={"Content-Type": "application/json",
                     "Origin": self.base})
        with urllib.request.urlopen(req, timeout=10) as resp:
            return json.load(resp)

    def _get_json(self, path: str) -> dict:
        with urllib.request.urlopen(self.base + path, timeout=10) as resp:
            return json.load(resp)

    def test_every_line_evidence_is_registered_in_facts(self):
        # 三家 × 两行等差清单：确认后 facts 证据注册表必须覆盖全部行证据
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["投标人名称", "统一社会信用代码", "清单编码",
                   "清单名称", "规格型号", "综合单价(元)"])
        for row in (["甲公司", "91330100MA1TEST00XY", "A1", "电线", "BV", 100],
                    ["乙公司", "91330100MA1TEST01AB", "A1", "电线", "BV", 110],
                    ["丙公司", "91330100MA1TEST02CD", "A1", "电线", "BV", 120],
                    ["甲公司", "91330100MA1TEST00XY", "A2", "开关", "10A", 200],
                    ["乙公司", "91330100MA1TEST01AB", "A2", "开关", "10A", 210],
                    ["丙公司", "91330100MA1TEST02CD", "A2", "开关", "10A", 220]):
            ws.append(row)
        buf = io.BytesIO()
        wb.save(buf)
        upload = self._post_json(
            "/api/upload?name=" + urllib.parse.quote("rect.xlsx"),
            {}) if False else None
        # 上传（二进制体）
        req = urllib.request.Request(
            self.base + "/api/upload?name=" + urllib.parse.quote("rect.xlsx"),
            data=buf.getvalue(), method="POST",
            headers={"Content-Type": "application/octet-stream",
                     "Origin": self.base})
        with urllib.request.urlopen(req, timeout=15) as resp:
            up = json.load(resp)
        self.assertTrue(up["ok"], up)

        state = self._get_json("/api/state")
        pl = [c for c in state["candidates"] if c["field"] == "price_lines"]
        self.assertEqual(len(pl), 3)
        bidder_rows = [c for c in state["candidates"]
                       if c["field"] == "bidder_name"]
        self.assertEqual(len(bidder_rows), 6)

        bidder_for = {"甲公司": "SYN-BIDDER-01", "乙公司": "SYN-BIDDER-02",
                      "丙公司": "SYN-BIDDER-03"}
        for group in pl:
            lines = group["value"] if isinstance(group["value"], list) \
                else json.loads(group["value"])
            # 以首行价格判断归属：100/110/120 对应 01/02/03
            price = lines[0]["unit_price"]
            bidder = {"100": "SYN-BIDDER-01", "110": "SYN-BIDDER-02",
                      "120": "SYN-BIDDER-03"}[str(price)]
            res = self._post_json("/api/confirm", {
                "event_id": "EV-SYN-001", "lot_id": "SYN-LOT-001",
                "bidder_id": bidder, "field": "price_lines",
                "value": lines, "evidence_id": group["evidence_id"],
                "action": "confirm"})
            self.assertTrue(res.get("ok"), res)

        facts = self.wb.build_p2_facts()
        registered = {e["evidence_id"] for e in facts.get("evidence", [])}
        missing = []
        for lot in facts["events"][0]["lots"]:
            for bid in lot["bids"]:
                for i, line in enumerate(bid.get("price_lines", []), start=1):
                    ev = line.get("evidence")
                    if ev and ev not in registered:
                        missing.append((bid["bidder_id"], i, ev))
                    for k, fe in (line.get("field_evidence") or {}).items():
                        if fe and fe not in registered:
                            missing.append((bid["bidder_id"], f"{i}:{k}", fe))
        self.assertEqual(missing, [],
                         "第 2 行起的清单行证据未注册进 facts 证据注册表："
                         f"{missing[:4]}")

        # 导入根因（复核五轮）：确认值里的字符串单价必须在此处归一为
        # 数值——事件导入曾因 unit_price 非数值整事件被拒（events=0）
        from biaojing import store
        import sqlite3 as _sqlite3
        p2_conn = store.connect(self.temp.name + "/p2_events.sqlite3")
        try:
            result = store.import_batch(facts, p2_conn)
            self.assertEqual(result.conflicts, [],
                             f"证据冲突：{result.conflicts}")
            self.assertEqual(result.events_inserted, 1,
                             "事件导入被拒，events_inserted 应为 1")
        finally:
            p2_conn.close()

    def test_decimal_precision_round_trip_through_store(self):
        # 0.29：float 最短 repr 可精确还原原十进制 → 确认→导入→读回无损
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["投标人名称", "综合单价(元)", "清单名称", "规格型号"])
        ws.append(["甲公司", 0.29, "电线", "BV"])
        buf = io.BytesIO()
        wb.save(buf)
        req = urllib.request.Request(
            self.base + "/api/upload?name=prec.xlsx", data=buf.getvalue(),
            method="POST",
            headers={"Content-Type": "application/octet-stream",
                     "Origin": self.base})
        with urllib.request.urlopen(req, timeout=15) as resp:
            self.assertTrue(json.load(resp)["ok"])
        state = self._get_json("/api/state")
        pl = next(c for c in state["candidates"]
                  if c["field"] == "price_lines")
        lines = pl["value"] if isinstance(pl["value"], list) \
            else json.loads(pl["value"])
        self.assertEqual(lines[0]["unit_price"], "0.29")
        res = self._post_json("/api/confirm", {
            "event_id": "EV-P", "lot_id": "LOT-P",
            "bidder_id": "SYN-BIDDER-01", "field": "price_lines",
            "value": lines, "evidence_id": pl["evidence_id"],
            "action": "confirm"})
        self.assertTrue(res.get("ok"), res)
        facts = self.wb.build_p2_facts()
        from decimal import Decimal
        line = facts["events"][0]["lots"][0]["bids"][0]["price_lines"][0]
        self.assertEqual(Decimal(str(line["unit_price"])), Decimal("0.29"))
        from biaojing import store
        p2_conn = store.connect(self.temp.name + "/p2_events.sqlite3")
        try:
            result = store.import_batch(facts, p2_conn)
            self.assertEqual(result.events_inserted, 1)
            stored = p2_conn.execute(
                "SELECT unit_price FROM bid_price_lines LIMIT 1").fetchone()[0]
            self.assertEqual(Decimal(str(stored)), Decimal("0.29"),
                             "REAL 读回经最短 repr 未还原 0.29，精度边界判断有误")
        finally:
            p2_conn.close()

    def test_overprecise_price_rejected_not_rounded(self):
        # 超出 float 可精确往返的金额：确认入口显式拒绝，绝不静默舍入
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["投标人名称", "综合单价(元)", "清单名称", "规格型号"])
        ws.append(["甲公司", 100, "电线", "BV"])
        buf = io.BytesIO()
        wb.save(buf)
        req = urllib.request.Request(
            self.base + "/api/upload?name=long.xlsx", data=buf.getvalue(),
            method="POST",
            headers={"Content-Type": "application/octet-stream",
                     "Origin": self.base})
        with urllib.request.urlopen(req, timeout=15) as resp:
            self.assertTrue(json.load(resp)["ok"])
        state = self._get_json("/api/state")
        pl = next(c for c in state["candidates"]
                  if c["field"] == "price_lines")
        lines = pl["value"] if isinstance(pl["value"], list) \
            else json.loads(pl["value"])
        lines[0]["unit_price"] = "0.123456789012345678901234"
        res = self._post_json("/api/confirm", {
            "event_id": "EV-Q", "lot_id": "LOT-Q",
            "bidder_id": "SYN-BIDDER-01", "field": "price_lines",
            "value": lines, "evidence_id": pl["evidence_id"],
            "action": "confirm"})
        self.assertFalse(res.get("ok"), "超精度单价被静默接受")
        self.assertNotIn("15", res.get("error", ""),
                         "错误提示不得写死位数上限（实际是逐值精确往返判定）")

    def test_non_numeric_price_returns_structured_error(self):
        # 复核七轮：'abc' 曾触发 decimal.InvalidOperation 未捕获 → 500。
        # 接口信任边界上的非法字符串必须返回结构化失败、确认数不变、
        # 服务不中断。
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["投标人名称", "综合单价(元)", "清单名称", "规格型号"])
        ws.append(["甲公司", 100, "电线", "BV"])
        buf = io.BytesIO()
        wb.save(buf)
        req = urllib.request.Request(
            self.base + "/api/upload?name=bad.xlsx", data=buf.getvalue(),
            method="POST",
            headers={"Content-Type": "application/octet-stream",
                     "Origin": self.base})
        with urllib.request.urlopen(req, timeout=15) as resp:
            self.assertTrue(json.load(resp)["ok"])
        state = self._get_json("/api/state")
        pl = next(c for c in state["candidates"]
                  if c["field"] == "price_lines")
        lines = pl["value"] if isinstance(pl["value"], list) \
            else json.loads(pl["value"])
        lines[0]["unit_price"] = "abc"
        try:
            res = self._post_json("/api/confirm", {
                "event_id": "EV-B", "lot_id": "LOT-B",
                "bidder_id": "SYN-BIDDER-01", "field": "price_lines",
                "value": lines, "evidence_id": pl["evidence_id"],
                "action": "confirm"})
        except urllib.error.HTTPError as exc:
            self.fail(f"非法单价触发 HTTP {exc.code}（应为结构化失败）")
        self.assertFalse(res.get("ok"), "非数值单价被静默接受")
        self.assertTrue(res.get("error"), "结构化失败须带可读 error")
        state2 = self._get_json("/api/state")
        self.assertEqual(state2.get("confirmation_count"), 0,
                         "被拒确认不得改变确认计数")
        # 服务不中断
        state3 = self._get_json("/api/state")
        self.assertEqual(state3["source_rows"][0]["status"], "success")


if __name__ == "__main__":
    unittest.main()
