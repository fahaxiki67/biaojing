# -*- coding: utf-8 -*-
"""标镜 P3 测试：候选抽取（真实 B01 合成样本）、工作台持久化与证据强制、
HTTP 安全边界、端到端筛查。B01 样本只读使用（synthetic）。

运行：
    cd app && python3 -W error::ResourceWarning -m unittest biaojing.tests.test_p3 -v
"""

from __future__ import annotations

import io
import json
import os
import sqlite3
import tempfile
import unittest
import urllib.request
from pathlib import Path
from unittest.mock import patch

import docx as docx_lib
import pymupdf

from biaojing import candidates, workspace as ws_mod
from biaojing.tests.test_p2 import ev, facts, bid, lot, event

B01 = Path(__file__).resolve().parents[3] / "samples"
B01_DOCX = B01 / "DOCX" / "SYN_B01_投标文件_蓝湾建设_一标段.docx"
B01_XLSX = B01 / "XLSX" / "SYN_B01_开标记录_一标段.xlsx"
B01_PDF = B01 / "PDF" / "SYN_B01_投标函_红枫机电_一标段_v1.pdf"
B01_AVAILABLE = B01_DOCX.exists() and B01_XLSX.exists() and B01_PDF.exists()


def load_p1(path: Path, parse):
    data = path.read_bytes()
    import hashlib
    sha = hashlib.sha256(data).hexdigest()
    rec = parse(data)
    tagged = candidates.assign_evidence_ids(sha, rec["evidence"])
    return sha, rec, tagged, data


class OcrEvidencePersistenceTests(unittest.TestCase):
    def test_ocr_marker_and_page_count_survive_workbench_ingest(self):
        from biaojing import pdf_parser

        binary, lang, _ = pdf_parser._ocr_runtime()
        if not binary or not lang:
            self.skipTest("本机缺少 Tesseract 或 chi_sim 中文模型")
        source = pymupdf.open()
        p = source.new_page(width=360, height=180)
        p.insert_text((32, 92), "OCR TEST 12345", fontsize=28,
                      fontname="helv")
        pix = p.get_pixmap(matrix=pymupdf.Matrix(200 / 72, 200 / 72),
                           alpha=False)
        scanned = pymupdf.open()
        scan_page = scanned.new_page(width=360, height=180)
        scan_page.insert_image(scan_page.rect, pixmap=pix)
        data = scanned.tobytes()
        source.close()
        scanned.close()

        wb = ws_mod.Workbench(tempfile.mkdtemp(prefix="biaojing_ocr_"))
        try:
            self.assertEqual(wb.ingest_bytes("scan.pdf", data)["status"],
                             "success")
            eid = wb.conn.execute(
                "SELECT evidence_id FROM evidence_store ORDER BY ordinal"
            ).fetchone()[0]
            evidence = wb.evidence_by_id(eid)
            self.assertTrue(evidence["ocr_used"])
            source_row = wb.state()["source_rows"][0]
            self.assertEqual(json.loads(source_row["counts_json"])["pages_ocr"], 1)
        finally:
            wb.close()

    def test_docx_image_ocr_candidate_and_locator_survive_workbench_ingest(self):
        from biaojing import pdf_parser

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.docx"
            data = pymupdf.Pixmap(
                pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), 0).tobytes("png")
            doc = docx_lib.Document()
            doc.add_paragraph().add_run().add_picture(io.BytesIO(data))
            doc.save(path)
            raw = path.read_bytes()
            wb = ws_mod.Workbench(directory)
            try:
                with patch.object(pdf_parser, "_ocr_runtime",
                                  return_value=("tesseract", "chi_sim+eng", None)), \
                        patch.object(pdf_parser, "ocr_image",
                                     return_value=("投标人名称：华夏建设有限公司", None)):
                    result = wb.ingest_bytes("投标文件.docx", raw)
                self.assertEqual(result["status"], "success")
                self.assertEqual(result["counts"]["images_ocr"], 1)
                row = wb.conn.execute(
                    "SELECT evidence_id FROM evidence_store WHERE kind='docx_image'")
                evidence = wb.evidence_by_id(row.fetchone()[0])
                self.assertEqual(evidence["locator_display"], "paragraph 1 image 1")
                self.assertEqual(evidence["page_status"], "ocr")
                self.assertTrue(evidence["ocr_used"])
                self.assertIn("embedded image OCR", evidence["extract_version"])
                candidate = wb.conn.execute(
                    "SELECT field, note, locator_json FROM candidates"
                    " WHERE evidence_id=?", (evidence["evidence_id"],)).fetchone()
                self.assertEqual(candidate["field"], "bidder_name")
                self.assertIn("人工复核", candidate["note"])
                self.assertEqual(json.loads(candidate["locator_json"])["kind"],
                                 "docx_image")
            finally:
                wb.close()


class FormulaEvidencePersistenceTests(unittest.TestCase):
    def test_xlsx_formula_evidence_keeps_cache_and_locator(self):
        import openpyxl
        from zipfile import ZIP_DEFLATED, ZipFile

        xlsx = io.BytesIO()
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "报价"
        sheet.append(["投标人", "综合单价(元)", "清单名称", "规格型号"])
        sheet.append(["甲公司", "=100+20", "电线", "BV"])
        book.save(xlsx)

        # openpyxl 不计算公式；仅在合成文件中写入模拟缓存，不接触真实资料。
        cached_xlsx = io.BytesIO()
        with ZipFile(io.BytesIO(xlsx.getvalue())) as source, \
                ZipFile(cached_xlsx, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    formula_xml = b"<f>100+20</f><v></v>"
                    self.assertIn(formula_xml, data)
                    data = data.replace(
                        formula_xml, b"<f>100+20</f><v>120</v>", 1)
                target.writestr(item, data)

        with tempfile.TemporaryDirectory(prefix="biaojing_formula_") as directory:
            workbench = ws_mod.Workbench(directory)
            try:
                result = workbench.ingest_bytes("报价.xlsx", cached_xlsx.getvalue())
                self.assertEqual(result["status"], "success")
                evidence_id = workbench.conn.execute(
                    "SELECT evidence_id FROM evidence_store"
                    " WHERE locator_display='报价!B2'"
                ).fetchone()[0]
                evidence = workbench.evidence_by_id(evidence_id)
                self.assertEqual(evidence["locator_display"], "报价!B2")
                self.assertIn("公式：=100+20", evidence["quote"])
                self.assertIn("缓存值：120", evidence["quote"])
                self.assertIn("可能尚未重新计算", evidence["quote"])
            finally:
                workbench.close()


# ---------------------------------------------------------------- 证据 ID

class EvidenceIdTests(unittest.TestCase):
    def test_deterministic_ids(self):
        evs = [{"kind": "docx_paragraph", "locator": "paragraph 1",
                "paragraph_index": 1, "text": "甲"},
               {"kind": "docx_paragraph", "locator": "paragraph 2",
                "paragraph_index": 2, "text": "乙"}]
        a = candidates.assign_evidence_ids("ab" * 32, evs)
        b = candidates.assign_evidence_ids("ab" * 32, evs)
        self.assertEqual([e["evidence_id"] for e in a],
                         [e["evidence_id"] for e in b])
        reordered = candidates.assign_evidence_ids("ab" * 32, list(reversed(evs)))
        self.assertEqual(a[0]["evidence_id"], reordered[1]["evidence_id"])
        self.assertTrue(a[0]["evidence_id"].startswith("E-" + "ab" * 32 + "-"))
        # 原条目不被修改（P1 locator/原文保持原样）
        self.assertNotIn("evidence_id", evs[0])
        self.assertEqual(evs[0]["locator"], "paragraph 1")

    def test_to_p2_locator_all_kinds(self):
        cases = [
            ({"kind": "pdf_page", "page": 2, "locator": "page 2"},
             {"kind": "pdf_page", "page": 2}),
            ({"kind": "docx_paragraph", "paragraph_index": 5},
             {"kind": "docx_paragraph", "paragraph_index": 5}),
            ({"kind": "docx_image", "image_index": 7},
             {"kind": "docx_image", "image_index": 7}),
            ({"kind": "docx_table_cell", "table_index": 1, "row": 2,
              "col": 3}, {"kind": "docx_table_cell", "table_index": 1,
                          "row": 2, "col": 3}),
            ({"kind": "xlsx_cell", "sheet": "S", "cell": "B2"},
             {"kind": "xlsx_cell", "sheet": "S", "cell": "B2"}),
            ({"kind": "pdf_page", "locator": "page 2"}, None),  # 缺 typed 字段
        ]
        for e, want in cases:
            self.assertEqual(candidates.to_p2_locator(e), want)


@unittest.skipUnless(B01_AVAILABLE, "B01 合成样本不在位")
class CandidateExtractionTests(unittest.TestCase):
    def test_docx_real_b01(self):
        sha, rec, tagged, _ = load_p1(B01_DOCX, __import__(
            "biaojing.docx_parser", fromlist=["parse"]).parse)
        cands = candidates.extract_candidates(sha, "docx", tagged)
        by_field = {}
        for c in cands:
            by_field.setdefault(c["field"], []).append(c)
        for field in ("project_name", "project_code", "lot_name",
                      "event_code", "bidder_name", "bidder_code",
                      "total_price", "contact_phone"):
            self.assertIn(field, by_field, f"缺字段候选 {field}")
        tp = by_field["total_price"][0]
        self.assertEqual(tp["value"]["status"], "normalized")
        self.assertEqual(tp["value"]["amount_yuan"], "119860000")
        # P2 locator 为 dict 且保留显示字符串
        self.assertIsInstance(by_field["project_name"][0]["locator"], dict)
        self.assertEqual(by_field["project_name"][0]["locator_display"],
                         "paragraph 3")

    def test_pdf_real_b01(self):
        sha, rec, tagged, _ = load_p1(B01_PDF, __import__(
            "biaojing.pdf_parser", fromlist=["parse"]).parse)
        cands = candidates.extract_candidates(sha, "pdf", tagged)
        fields = {c["field"] for c in cands}
        self.assertIn("bidder_name", fields)
        self.assertIn("total_price", fields)

    def test_xlsx_header_suggestions_real_b01(self):
        sha, rec, tagged, _ = load_p1(B01_XLSX, __import__(
            "biaojing.xlsx_parser", fromlist=["parse"]).parse)
        cands = candidates.extract_candidates(sha, "xlsx", tagged)
        by_field = {}
        for c in cands:
            by_field.setdefault(c["field"], []).append(c)
        # 开标记录表头：投标单位/总报价(元) 列建议
        self.assertGreaterEqual(len(by_field.get("bidder_name", [])), 3)
        prices = [c for c in by_field.get("total_price", [])
                  if "表头" in c["note"]]
        self.assertGreaterEqual(len(prices), 3)
        self.assertEqual(prices[0]["locator"]["kind"], "xlsx_cell")
        # 公式串不给候选
        self.assertFalse(any(str(c["value"]).startswith("=")
                             for c in cands))


# ---------------------------------------------------------------- 工作台

@unittest.skipUnless(B01_AVAILABLE, "B01 合成样本不在位")
class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="biaojing_p3_")
        self.wb = ws_mod.Workbench(self.tmp)

    def tearDown(self):
        self.wb.close()

    def test_ingest_coverage_duplicate_and_hash_path(self):
        data = B01_DOCX.read_bytes()
        r1 = self.wb.ingest_bytes("蓝湾/投标文件.docx", data)
        self.assertEqual(r1["status"], "success")
        r2 = self.wb.ingest_bytes("副本.docx", data)
        self.assertEqual(r2["status"], "duplicate")
        cov = self.wb.coverage()
        # 覆盖率按输入出现次数：两次上传 = 2 次出现（1 success + 1 duplicate）
        self.assertEqual(cov["total_input_occurrences"], 2)
        self.assertEqual(cov["by_status"]["success"], 1)
        self.assertEqual(cov["by_status"]["duplicate"], 1)
        self.assertEqual(cov["unique_files"], 1)
        # 字节按哈希路径保存，客户端文件名不作为存储路径
        fp = os.path.join(self.tmp, "files", r1["sha256"][:2], r1["sha256"])
        self.assertTrue(os.path.exists(fp))
        self.assertEqual(Path(fp).read_bytes(), data)
        # 同源重复导入 → evidence ID 确定性一致
        cands = self.wb.conn.execute(
            "SELECT evidence_id FROM candidates ORDER BY id").fetchall()
        self.assertTrue(all(r["evidence_id"].startswith(
            "E-" + r1["sha256"][:12]) for r in cands))

    def test_state_paginates_source_occurrences_and_returns_compact_summary(self):
        self.wb.conn.executemany(
            "INSERT INTO source_refs(ref,status,seen_at) VALUES(?,?,?)",
            [(f"source-{i}.pdf", "unknown", "2026-01-01")
             for i in range(3)])
        self.wb.conn.commit()

        first = self.wb.state(source_limit=2)
        second = self.wb.state(source_offset=2, source_limit=2)
        self.assertEqual([r["ref"] for r in first["source_rows"]],
                         ["source-0.pdf", "source-1.pdf"])
        self.assertEqual(first["source_total"], 3)
        self.assertTrue(first["source_has_more"])
        self.assertEqual([r["ref"] for r in second["source_rows"]],
                         ["source-2.pdf"])
        self.assertFalse(second["source_has_more"])
        self.assertEqual(second["confirmation_count"], 0)
        self.assertNotIn("sources", second)
        self.assertNotIn("confirmations", second)
        self.assertNotIn("file_bindings", second)

    def test_xlsx_extraction_version_is_recorded(self):
        from biaojing import xlsx_parser

        result = self.wb.ingest_bytes("开标记录.xlsx", B01_XLSX.read_bytes())
        self.assertEqual(result["status"], "success")
        source = self.wb.conn.execute(
            "SELECT extract_version FROM sources WHERE sha256=?",
            (result["sha256"],)).fetchone()
        self.assertEqual(source["extract_version"],
                         xlsx_parser.EXTRACTION_VERSION)

    def test_confirm_requires_resolvable_evidence(self):
        data = B01_DOCX.read_bytes()
        self.wb.ingest_bytes("蓝湾.docx", data)
        ok = self.wb.confirm_field("EV-1", "L1", "B1", "total_price",
                                   100.0, "E-ghost", "confirm")
        self.assertFalse(ok["ok"])
        self.assertIn("不得确认", ok["error"])

    def test_e2e_confirm_screen_finding_and_registry(self):
        # 验收 4：同 event+lot 两个 B01 投标人共享联系电话（蓝湾 DOCX 与
        # 红枫 PDF 均为 138-0000-0001）→ 确认后 R002 触发，finding 引用
        # 双方精确来源证据，全部经证据端点可解析
        sha_docx, _, _, _ = load_p1(B01_DOCX, __import__(
            "biaojing.docx_parser", fromlist=["parse"]).parse)
        r1 = self.wb.ingest_bytes("蓝湾.docx", B01_DOCX.read_bytes())
        r2 = self.wb.ingest_bytes("红枫.pdf", B01_PDF.read_bytes())
        cands = self.wb.conn.execute(
            "SELECT sha256, field, value_json, evidence_id FROM candidates"
            " WHERE field='contact_phone' ORDER BY id").fetchall()
        shared = [c for c in cands
                  if "138-0000-0001" in str(json.loads(c["value_json"]))]
        self.assertEqual(len(shared), 2, "两来源应各有共享电话候选")
        phone1, phone2 = shared
        self.assertNotEqual(phone1["sha256"], phone2["sha256"])
        for cand, bidder in ((phone1, "SYN-BIDDER-01"), (phone2, "SYN-BIDDER-02")):
            ok = self.wb.confirm_field(
                "EV-2026-001-01", "SYN-LOT-001", bidder, "contact_phone",
                json.loads(cand["value_json"]), cand["evidence_id"],
                "confirm", source_role="bidder")
            self.assertTrue(ok["ok"], ok.get("error"))
        result = self.wb.run_screen()
        r002 = [f for f in result["findings"] if f["rule_id"] == "R002"]
        self.assertEqual(len(r002), 1)
        self.assertEqual(r002[0]["scope"]["bidder_ids"],
                         ["SYN-BIDDER-01", "SYN-BIDDER-02"])
        expected_evs = {phone1["evidence_id"], phone2["evidence_id"]}
        self.assertTrue(expected_evs <= set(r002[0]["evidence_ids"]))
        for f in result["findings"]:
            for eid in f["evidence_ids"]:
                self.assertIsNotNone(self.wb.evidence_by_id(eid), eid)
        self.assertEqual(result["unresolved_evidence"], [])
        self.assertEqual(result["import"]["event_count_after"], 1)

        # 验收 5：把一方电话标 unknown 后重筛 → R002 随事实消失
        self.wb.confirm_field("EV-2026-001-01", "SYN-LOT-001",
                              "SYN-BIDDER-02", "contact_phone",
                              "unknown", None, "unknown")
        result2 = self.wb.run_screen()
        self.assertEqual([f for f in result2["findings"]
                          if f["rule_id"] == "R002"], [])
        # 重复筛查：事件次数不变（幂等累计）
        self.assertEqual(result2["import"]["event_count_after"], 1)

    def test_r002_contact_role_gate_default_unknown(self):
        # 来源角色门控：确认电话数值 ≠ 确认来源角色。默认 unknown 不触发；
        # 显式选 bidder 才触发；一方 agency 一方 bidder 也不触发
        shared = [("蓝湾.docx", B01_DOCX, "SYN-BIDDER-01"),
                  ("红枫.pdf", B01_PDF, "SYN-BIDDER-02")]
        for roles, want_finding in (
                ((None, None), False),            # 双方默认 unknown
                (("agency", "bidder"), False),    # 一方代理
                (("bidder", "bidder"), True)):    # 双方显式 bidder
            self.wb2 = ws_mod.Workbench(tempfile.mkdtemp(prefix="biaojing_p3_role_"))
            try:
                for (fname, path, bidder), role in zip(shared, roles):
                    self.wb2.ingest_bytes(fname, path.read_bytes())
                cands = self.wb2.conn.execute(
                    "SELECT sha256, field, value_json, evidence_id FROM"
                    " candidates WHERE field='contact_phone'"
                    " ORDER BY id").fetchall()
                shared_phones = [c for c in cands
                                 if "138-0000-0001" in str(
                                     json.loads(c["value_json"]))]
                self.assertEqual(len(shared_phones), 2)
                for cand, bidder, role in zip(shared_phones,
                                              ("SYN-BIDDER-01",
                                               "SYN-BIDDER-02"), roles):
                    ok = self.wb2.confirm_field(
                        "EV-2026-001-01", "SYN-LOT-001", bidder,
                        "contact_phone", json.loads(cand["value_json"]),
                        cand["evidence_id"], "confirm",
                        source_role=role or "unknown")
                    self.assertTrue(ok["ok"])
                result = self.wb2.run_screen()
                got = [f for f in result["findings"]
                       if f["rule_id"] == "R002"]
                self.assertEqual(bool(got), want_finding,
                                 f"roles={roles} 时 R002 应为 {want_finding}")
            finally:
                self.wb2.close()

    def test_old_schema_workspace_migration(self):
        # 旧版工作区（candidates 无 companion_json、confirmations 无
        # source_role）打开后自动补列，既有数据继续可用
        oldws = os.path.join(self.tmp, "oldws")
        os.makedirs(oldws)
        db = os.path.join(oldws, "biaojing.sqlite3")
        conn = sqlite3.connect(db)
        conn.executescript(
            "CREATE TABLE IF NOT EXISTS candidates("
            " id INTEGER PRIMARY KEY AUTOINCREMENT, sha256 TEXT, field TEXT,"
            " value_json TEXT, evidence_id TEXT, locator_json TEXT,"
            " locator_display TEXT, note TEXT);"
            "CREATE TABLE IF NOT EXISTS confirmations("
            " id INTEGER PRIMARY KEY AUTOINCREMENT,"
            " event_id TEXT NOT NULL, lot_id TEXT NOT NULL,"
            " bidder_id TEXT NOT NULL, field TEXT NOT NULL, value TEXT,"
            " evidence_id TEXT, original_candidate TEXT, action TEXT,"
            " confirmed_at TEXT,"
            " UNIQUE(event_id, lot_id, bidder_id, field));"
            # 预置一行旧候选与一条旧确认，验证迁移后数据仍在
            "INSERT INTO candidates(sha256, field, value_json, evidence_id,"
            " locator_json, locator_display, note) VALUES('oldsha',"
            " 'total_price', '119860000.0', 'E-oldsha-0001', '{}',"
            " 'paragraph 8', '标签匹配');"
            "INSERT INTO confirmations(event_id, lot_id, bidder_id, field,"
            " value, evidence_id, original_candidate, action, confirmed_at)"
            " VALUES('EV-OLD', 'L1', 'B1', 'bidder_name', '旧名称',"
            " 'E-oldsha-0001', NULL, 'confirm', '2026-01-01');")
        conn.commit()
        conn.close()
        wb = ws_mod.Workbench(oldws)  # 直接以旧库路径打开
        try:
            cols = {r[1] for r in wb.conn.execute(
                "PRAGMA table_info(candidates)")}
            self.assertIn("companion_json", cols)
            ccols = {r[1] for r in wb.conn.execute(
                "PRAGMA table_info(confirmations)")}
            self.assertIn("source_role", ccols)
            # 迁移后既有数据仍在（旧行自动获得默认值）
            kept = wb.conn.execute(
                "SELECT value, source_role FROM confirmations"
                " WHERE event_id='EV-OLD'").fetchone()
            self.assertEqual(kept["value"], "旧名称")
            self.assertEqual(kept["source_role"] or "unknown", "unknown")
            kept_cand = wb.conn.execute(
                "SELECT value_json, companion_json FROM candidates"
                " WHERE evidence_id='E-oldsha-0001'").fetchone()
            self.assertEqual(json.loads(kept_cand["value_json"]), 119860000.0)
            self.assertIsNone(kept_cand["companion_json"])
            # 旧库上正常写入候选与确认（含新列）
            wb.ingest_bytes("蓝湾.docx", B01_DOCX.read_bytes())
            cands = wb.conn.execute(
                "SELECT evidence_id FROM candidates"
                " WHERE field='total_price' AND evidence_id != 'E-oldsha-0001'"
                " LIMIT 1").fetchall()
            ok = wb.confirm_amount("EV-1", "L1", "SYN-BIDDER-01",
                                   "119860000", "yuan", "CNY", True,
                                   cands[0]["evidence_id"], "confirm")
            self.assertTrue(ok["ok"])
        finally:
            wb.close()

    def test_unknown_and_correction_flow(self):
        data = B01_DOCX.read_bytes()
        self.wb.ingest_bytes("蓝湾.docx", data)
        # 标 unknown：无需证据
        ok = self.wb.confirm_field("EV-1", "L1", "B1", "uscc",
                                   "任意", None, "unknown")
        self.assertTrue(ok["ok"])
        row = self.wb.conn.execute(
            "SELECT value, action FROM confirmations WHERE field='uscc'"
        ).fetchone()
        self.assertEqual(row["value"], "unknown")
        self.assertEqual(row["action"], "unknown")
        # 更正：保留原始候选值
        cands = self.wb.conn.execute(
            "SELECT field, value_json, evidence_id FROM candidates"
            " WHERE field='total_price' LIMIT 1").fetchall()
        tp = cands[0]
        self.wb.confirm_amount("EV-1", "L1", "B1", "55", "yuan", "CNY",
                               "unknown", tp["evidence_id"], "correct",
                               original_candidate=tp["value_json"])
        row = self.wb.conn.execute(
            "SELECT value, original_candidate, action FROM confirmations"
            " WHERE field='total_price'").fetchone()
        self.assertEqual(json.loads(row["value"]), "55")
        # 原始结构化候选被原样保留，而非更正值。
        self.assertEqual(json.loads(row["original_candidate"])["amount_yuan"],
                         "119860000")
        self.assertEqual(row["action"], "correct")

    def test_cny_amount_roundtrip_and_tax_default_unknown(self):
        # 原始人民币金额与单位进入同一确认事务；税口径 unknown 不进 R005 分母。
        data = B01_DOCX.read_bytes()
        self.wb.ingest_bytes("蓝湾.docx", data)
        cands = self.wb.conn.execute(
            "SELECT field, value_json, evidence_id, candidate_meta_json FROM"
            " candidates WHERE field='total_price'"
            " ORDER BY id").fetchall()
        tp = cands[0]
        cand_value = json.loads(tp["value_json"])
        meta = json.loads(tp["candidate_meta_json"])
        self.assertEqual(meta["currency_hint"], "CNY")
        self.wb.confirm_amount("EV-1", "L1", "SYN-BIDDER-01",
                               cand_value["raw"], cand_value["unit_hint"],
                               meta["currency_hint"], "unknown",
                               tp["evidence_id"], "confirm")
        result = self.wb.run_screen()
        r005 = [f for f in result["findings"] if f["rule_id"] == "R005"]
        # 税口径 unknown → 该报价被隔离，R005 不触发集中度
        self.assertTrue(all(f["signal"] == "不计算" for f in r005))
        # roundtrip：确认值与证据经库读回保持
        facts = self.wb.build_p2_facts()
        b1 = facts["events"][0]["lots"][0]["bids"][0]
        self.assertEqual(b1["total_price"], "119860000")
        self.assertIsNone(b1.get("tax_included"))
        self.assertIsNotNone(b1["total_price_evidence"])

    def test_multi_contacts_people_file_binding_and_screen_snapshot(self):
        import hashlib

        self.wb.ingest_bytes("蓝湾.docx", B01_DOCX.read_bytes())
        sha = hashlib.sha256(B01_DOCX.read_bytes()).hexdigest()
        evidence_ids = [r[0] for r in self.wb.conn.execute(
            "SELECT evidence_id FROM evidence_store WHERE sha256=?"
            " ORDER BY ordinal LIMIT 4", (sha,))]
        self.assertEqual(len(evidence_ids), 4)
        for value, eid in zip(("138-0000-0001", "138-0000-0002"), evidence_ids[:2]):
            self.assertTrue(self.wb.confirm_field(
                "EV-AUDIT", "LOT-AUDIT", "BID-A", "contact_phone", value,
                eid, "confirm", source_role="bidder")["ok"])
        for name, person_id, eid in (
                ("张三", "PERSON-A", evidence_ids[2]),
                ("李四", "PERSON-B", evidence_ids[3])):
            self.assertTrue(self.wb.confirm_field(
                "EV-AUDIT", "LOT-AUDIT", "BID-A", "person_manager", name,
                eid, "confirm", person_id=person_id)["ok"])
        self.assertTrue(self.wb.confirm_field(
            "EV-AUDIT", "LOT-AUDIT", "BID-B", "person_tech", "张三",
            evidence_ids[2], "confirm", person_id="PERSON-A")["ok"])
        self.assertTrue(self.wb.confirm_field(
            "EV-AUDIT", "LOT-AUDIT", "BID-A", "outcome_result", "未中标",
            evidence_ids[3], "confirm")["ok"])
        self.assertTrue(self.wb.bind_file(
            "EV-AUDIT", "LOT-AUDIT", "BID-A", sha, "bid_document",
            "bid_document", declared_owner_id="BID-OTHER")["ok"])

        facts = self.wb.build_p2_facts()
        stored_bid = facts["events"][0]["lots"][0]["bids"][0]
        second_bid = facts["events"][0]["lots"][0]["bids"][1]
        self.assertEqual(len(stored_bid["contacts"]), 2)
        self.assertEqual({p["person_id"] for p in stored_bid["persons"]},
                         {"PERSON-A", "PERSON-B"})
        self.assertEqual(len(stored_bid["files"]), 1)
        self.assertEqual(stored_bid["outcome"]["result"], "lost")
        self.assertEqual(stored_bid["outcome"]["evidence"], evidence_ids[3])
        self.assertEqual(second_bid["persons"][0]["person_id"], "PERSON-A")
        screened = self.wb.run_screen()
        self.assertTrue(any(f["rule_id"] == "R001"
                            for f in screened["findings"]))
        self.assertTrue(any(f["rule_id"] == "R003"
                            for f in screened["findings"]))
        run = self.wb.conn.execute(
            "SELECT facts_sha256,facts_json FROM screen_runs"
            " ORDER BY id DESC LIMIT 1").fetchone()
        snapshot = json.loads(run["facts_json"])
        self.assertEqual(run["facts_sha256"], hashlib.sha256(
            json.dumps(snapshot, ensure_ascii=False, sort_keys=True,
                       separators=(",", ":")).encode("utf-8")).hexdigest())
        state = self.wb.state()
        self.assertEqual(state["file_binding_history_total"], 1)
        self.assertEqual(state["confirmation_history_total"], 6)
        self.assertEqual(state["file_binding_history"][0]["evidence_id"],
                         evidence_ids[0])
        self.assertEqual(state["last_screen_run"]["facts_sha256"],
                         run["facts_sha256"])

    def test_empty_and_failed_uploads(self):
        result = self.wb.run_screen()  # 空工作区：0 finding，不崩
        self.assertEqual(result["summary"]["total"], 0)
        self.assertEqual(self.wb.run_screen()["import"]["event_count_after"], 0)
        r = self.wb.ingest_bytes("坏文件.docx", b"not a docx at all")
        self.assertEqual(r["status"], "failed")
        cov = self.wb.coverage()
        self.assertEqual(cov["by_status"]["failed"], 1)

    def test_first_reason_empty_repeat_duplicate_reason(self):
        # 独立复核：首次上传 reason 必须为空（无重复原因），重复 SHA
        # 才是 duplicate 且带"与先前来源字节相同"——不得写反
        data = B01_DOCX.read_bytes()
        r1 = self.wb.ingest_bytes("首次.docx", data)
        self.assertEqual(r1["status"], "success")
        row1 = self.wb.conn.execute(
            "SELECT status, reason FROM source_refs WHERE sha256=?"
            " ORDER BY id", (r1["sha256"],)).fetchone()
        self.assertIsNone(row1["reason"])
        r2 = self.wb.ingest_bytes("重复.docx", data)
        self.assertEqual(r2["status"], "duplicate")
        row2 = self.wb.conn.execute(
            "SELECT status, reason FROM source_refs WHERE sha256=?"
            " ORDER BY id DESC LIMIT 1", (r2["sha256"],)).fetchone()
        self.assertEqual(row2["status"], "duplicate")
        self.assertEqual(row2["reason"], "与先前来源字节相同")

    def test_zip_roundtrip_hash_and_coverage(self):
        # ZIP 容器本体按哈希存证 + 归档级 occurrence；首次 archived 恒保留；
        # 展开条目（2 有效 + 1 拒收穿越）分别记录；同 SHA 重复归档 →
        # duplicate（+3）且条目仍展开，拒收再 +1，合计 +4 次出现
        import zipfile

        archive = os.path.join(self.tmp, "pack.zip")
        with zipfile.ZipFile(archive, "w") as zf:
            zf.writestr("docs/a.docx", B01_DOCX.read_bytes())
            zf.writestr("docs/b.pdf", B01_PDF.read_bytes())
            zf.writestr("../逃逸.txt", "x")
        archive_bytes = Path(archive).read_bytes()
        results = self.wb.ingest_zip("pack.zip", archive_bytes)
        # 归档本体已按哈希路径保留且字节一致
        import hashlib
        sha = hashlib.sha256(archive_bytes).hexdigest()
        stored = os.path.join(self.tmp, "files", sha[:2], sha)
        self.assertTrue(os.path.exists(stored))
        self.assertEqual(Path(stored).read_bytes(), archive_bytes)
        cov = self.wb.coverage()
        self.assertEqual(cov["by_status"]["archived"], 1)
        self.assertEqual(cov["by_status"]["rejected"], 1)  # 穿越条目
        self.assertTrue(any(r["status"] == "success" for r in results))
        # 同 SHA 重复归档 → duplicate，但条目仍展开（各成 duplicate）：
        # 新增 4 次出现 = 归档 duplicate + 2 条目 duplicate + 1 拒收重复
        results2 = self.wb.ingest_zip("pack.zip", archive_bytes)
        self.assertEqual(results2[0]["status"], "duplicate")
        cov2 = self.wb.coverage()
        self.assertEqual(cov2["by_status"]["archived"], 1)  # 首次归档行保留
        self.assertEqual(cov2["by_status"]["duplicate"],
                         cov["by_status"].get("duplicate", 0) + 3)
        self.assertEqual(cov2["by_status"]["rejected"],
                         cov["by_status"]["rejected"] + 1)
        self.assertEqual(cov2["total_input_occurrences"],
                         cov["total_input_occurrences"] + 4)
        # 展开条目各自可解析
        self.assertTrue(any(r["status"] == "success" for r in results))

    def test_zip_expand_failure_marks_failed(self):
        # 展开失败：归档 occurrence 原地转 failed（不留 archived+failed 双行）
        archive = os.path.join(self.tmp, "bad.zip")
        fake = b"PK\x03\x04" + b"\x00" * 64  # PK 头但结构损坏
        Path(archive).write_bytes(fake)
        results = self.wb.ingest_zip("bad.zip", fake)
        self.assertEqual(results[-1]["status"], "failed")
        cov = self.wb.coverage()
        self.assertEqual(cov["by_status"].get("archived"), None)
        self.assertEqual(cov["by_status"]["failed"], 1)
        import hashlib
        sha = hashlib.sha256(fake).hexdigest()
        statuses = [r["status"] for r in self.wb.conn.execute(
            "SELECT status FROM source_refs WHERE sha256=?", (sha,))]
        self.assertEqual(statuses, ["failed"])


# ---------------------------------------------------------------- HTTP 安全 + 冒烟

class HttpSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="biaojing_p3_http_")
        from biaojing import webapp
        cls.server = webapp.WorkbenchServer(
            ("127.0.0.1", 0), ws_mod.Workbench(cls.tmp))
        cls.port = cls.server.server_address[1]
        import threading
        cls.thread = threading.Thread(target=cls.server.serve_forever,
                                      daemon=True)
        cls.thread.start()
        cls.base = f"http://127.0.0.1:{cls.port}"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.server.workbench.close()  # 严格模式下 SQLite 连接必须关闭

    def get(self, path, headers=None):
        req = urllib.request.Request(self.base + path, headers=headers or {})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            try:
                return e.code, e.read()
            finally:
                e.close()

    def post(self, path, body: bytes, headers=None, origin: str | None = "auto"):
        h = {"Content-Type": "application/octet-stream"}
        if origin == "auto":
            h["Origin"] = self.base
        elif origin is not None:
            h["Origin"] = origin
        h.update(headers or {})
        req = urllib.request.Request(self.base + path, data=body,
                                     headers=h, method="POST")
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            try:
                return e.code, e.read()
            finally:
                e.close()

    def test_home_and_state(self):
        code, body = self.get("/")
        self.assertEqual(code, 200)
        self.assertIn("标镜".encode(), body)
        self.assertIn("aboutButton".encode(), body)
        code, body = self.get("/api/state")
        self.assertEqual(code, 200)
        self.assertEqual(json.loads(body)["product_name"], "标镜")
        code, body = self.get("/api/about")
        self.assertEqual(code, 200)
        about = json.loads(body)
        self.assertEqual(about["author"], "刘奇")
        from biaojing import VERSION
        self.assertEqual(about["version"], VERSION)

    def test_update_check_is_explicitly_unconfigured_without_repository(self):
        from unittest.mock import patch
        with patch.dict(os.environ, {"BIAOJING_GITHUB_REPOSITORY": ""}):
            code, body = self.post("/api/update/check", b"{}",
                                   {"Content-Type": "application/json"})
        self.assertEqual(code, 200)
        self.assertEqual(json.loads(body)["status"], "unconfigured")

    def test_missing_origin_rejected(self):
        # 二次复核补充：mutating POST 缺 Origin 一律拒绝
        code, _ = self.post("/api/screen", b"{}", origin=None)
        self.assertEqual(code, 403)

    def test_prefix_lookalike_host_and_origin_rejected(self):
        # 前缀相似域与错误端口都不得通过精确同源校验
        code, _ = self.get("/api/state", {"Host": "127.0.0.1.evil.com"})
        self.assertEqual(code, 403)
        code, _ = self.post("/api/screen", b"{}",
                            origin=f"http://127.0.0.1.evil.com:{self.port}")
        self.assertEqual(code, 403)
        code, _ = self.post("/api/screen", b"{}",
                            origin=f"http://127.0.0.1:{self.port + 1}")
        self.assertEqual(code, 403)

    def test_oversized_body_rejected_and_recorded(self):
        # 把上限压到 1 字节：2 字节请求体证明"读取前拒绝"（不发大文件）
        from biaojing import webapp
        real_cap = webapp.MAX_BODY
        webapp.MAX_BODY = 1
        try:
            code, body = self.post("/api/upload?name=oversize.bin", b"xx",
                                   {"Content-Length": "2"})
            self.assertEqual(code, 413)
            self.assertIn("超过上限".encode(), body)  # 中文详情在 UTF-8 JSON 体
        finally:
            webapp.MAX_BODY = real_cap
        state = json.loads(self.get("/api/state")[1])
        rej = [r for r in state["source_rows"] if r["status"] == "rejected"]
        self.assertTrue(rej and "超过上限" in rej[0]["reason"])

    def test_reject_record_endpoint_200(self):
        # HTTP 层拒收上报端点：同源 POST 返回 200，且记录进入覆盖率
        code, body = self.post("/api/reject_record", json.dumps({
            "name": "枚举失败目录", "reason": "读取失败：权限"}).encode(),
            {"Content-Type": "application/json"})
        self.assertEqual(code, 200)
        self.assertTrue(json.loads(body)["ok"])
        state = json.loads(self.get("/api/state")[1])
        rej = [r for r in state["source_rows"] if r["status"] == "rejected"]
        self.assertTrue(any("读取失败" in (r["reason"] or "") for r in rej))

    def test_download_percent2f_name_roundtrip(self):
        # HTTP 级：文件名含字面 %2F（上传时经 quote 双重编码），下载返回
        # 200 + nosniff + filename*=UTF-8''，body 逐字节等于原文件
        from urllib.parse import quote
        name = "含%2F百分号_蓝湾.docx"  # 字面 %2F 序列
        data = B01_DOCX.read_bytes()
        code, _ = self.post(f"/api/upload?name={quote(name, safe='')}", data)
        self.assertEqual(code, 200)
        state = json.loads(self.get("/api/state")[1])
        row = [r for r in state["source_rows"]
               if "含%2F百分号" in (r["ref"] or "")]
        self.assertTrue(row, "来源引用须保留字面 %2F，不得改写为路径分隔符")
        sha = state["candidates"][0]["sha256"]
        req = urllib.request.Request(f"{self.base}/api/download/{sha}",
                                     headers={"Origin": self.base})
        with urllib.request.urlopen(req) as r:
            self.assertEqual(r.status, 200)
            self.assertEqual(r.headers.get("X-Content-Type-Options"), "nosniff")
            cd = r.headers.get("Content-Disposition", "")
            self.assertIn("attachment", cd)
            self.assertIn("filename*=UTF-8''", cd)
            self.assertEqual(r.read(), data)  # 逐字节等于已存原件
        # 非法 SHA（长度不足 / 非 hex / 路径穿越）一律拒绝
        for bad in ("abc", "z" * 64, "../" + "a" * 64):
            ecode, _ = self.get("/api/download/" + urllib.parse.quote(bad))
            self.assertIn(ecode, (400, 404))

    def test_upload_screen_roundtrip(self):
        from biaojing import webapp
        from urllib.parse import quote
        real_cap = webapp.MAX_BODY
        webapp.MAX_BODY = ws_mod.MAX_UPLOAD_BYTES  # 本测试用真实上限
        try:
            state = json.loads(self.get("/api/state")[1])
            prior = state["coverage"]["total_input_occurrences"]
            data = B01_DOCX.read_bytes()
            code, _ = self.post(
                f"/api/upload?name={quote(os.path.basename(str(B01_DOCX)))}",
                data)
            self.assertEqual(code, 200)
            state = json.loads(self.get("/api/state")[1])
            self.assertEqual(state["coverage"]["total_input_occurrences"],
                             prior + 1)
            # 确认共享电话（两来源，角色显式 bidder）→ R002 经 HTTP 栈触发
            conf = {}
            self.post(f"/api/upload?name={quote(B01_PDF.name)}",
                      B01_PDF.read_bytes())
            # 分页接口返回稳定候选页，不把全量数据塞进首次界面状态。
            page1 = json.loads(self.get(
                "/api/state?candidate_offset=0&candidate_limit=1")[1])
            self.assertEqual(len(page1["candidates"]), 1)
            self.assertGreater(page1["candidate_total"], 1)
            self.assertTrue(page1["candidate_has_more"])
            page2 = json.loads(self.get(
                "/api/state?candidate_offset=1&candidate_limit=1")[1])
            self.assertEqual(len(page2["candidates"]), 1)
            self.assertNotEqual(page1["candidates"][0]["id"],
                                page2["candidates"][0]["id"])
            sha_docx = __import__("hashlib").sha256(
                B01_DOCX.read_bytes()).hexdigest()
            sha_pdf = __import__("hashlib").sha256(
                B01_PDF.read_bytes()).hexdigest()
            cands = json.loads(self.get("/api/state")[1])["candidates"]
            for sha_prefix, bidder in ((sha_docx, "SYN-BIDDER-01"),
                                       (sha_pdf, "SYN-BIDDER-02")):
                phone = [c for c in cands if c["field"] == "contact_phone"
                         and "138-0000-0001" in str(c["value"])
                         and c["evidence_id"].startswith("E-" + sha_prefix[:12])]
                self.assertTrue(phone, f"{bidder} 缺共享电话候选")
                conf[bidder] = phone[0]
            for bidder, cand in conf.items():
                code, body = self.post(
                    "/api/confirm", json.dumps({
                        "event_id": "EV-2026-001-01", "lot_id": "SYN-LOT-001",
                        "bidder_id": bidder, "field": "contact_phone",
                        "value": cand["value"],
                        "evidence_id": cand["evidence_id"],
                        "source_role": "bidder",
                        "action": "confirm"}).encode(),
                    {"Content-Type": "application/json"})
                self.assertEqual(code, 200)
                self.assertTrue(json.loads(body)["ok"])
            amount = next(c for c in json.loads(self.get("/api/state")[1])["candidates"]
                          if c["field"] == "total_price"
                          and c["sha256"] == sha_docx)
            value = amount["value"]
            code, body = self.post("/api/confirm_amount", json.dumps({
                "event_id": "EV-2026-001-01", "lot_id": "SYN-LOT-001",
                "bidder_id": "SYN-BIDDER-01", "raw_value": value["raw"],
                "unit": value["unit_hint"],
                "currency": amount.get("currency_hint") or "unknown",
                "tax_included": "unknown",
                "evidence_id": amount["evidence_id"], "action": "confirm",
            }).encode(), {"Content-Type": "application/json"})
            self.assertEqual(code, 200)
            self.assertTrue(json.loads(body)["ok"])
            code, body = self.post("/api/bind_file", json.dumps({
                "event_id": "EV-2026-001-01", "lot_id": "SYN-LOT-001",
                "bidder_id": "SYN-BIDDER-01", "sha256": sha_docx,
                "context": "bid_document", "source_type": "bid_document",
                "declared_owner_id": "SYN-BIDDER-OTHER",
            }).encode(), {"Content-Type": "application/json"})
            self.assertEqual(code, 200)
            self.assertTrue(json.loads(body)["ok"])
            code, body = self.post("/api/screen", b"{}")
            self.assertEqual(code, 200)
            payload = json.loads(body)
            r002 = [f for f in payload["findings"] if f["rule_id"] == "R002"]
            self.assertEqual(len(r002), 1)
            self.assertTrue(any(f["rule_id"] == "R001"
                                for f in payload["findings"]))
            refreshed = json.loads(self.get("/api/state")[1])
            self.assertEqual(refreshed["last_screen_run"]["status"], "completed")
            self.assertEqual(len(refreshed["last_screen_run"]["facts_sha256"]), 64)
            self.assertEqual(refreshed["file_binding_history_total"], 1)
            for f in payload["findings"]:
                for eid in f["evidence_ids"]:
                    ecode, ebody = self.get("/api/evidence/" + eid)
                    self.assertEqual(ecode, 200)
        finally:
            webapp.MAX_BODY = real_cap


# ---------------------------------------------------------------- UI 渲染安全

class UiRenderSafetyTests(unittest.TestCase):
    """二次复核补充：文档抽取文本全部经 textContent 渲染，
    不允许 innerHTML 插值 / 内联事件处理器 / document.write。"""

    def test_ui_page_has_no_html_injection_channels(self):
        from biaojing import ui_page

        src = ui_page.INDEX_HTML
        for banned in ("innerHTML", "outerHTML", "document.write",
                       "insertAdjacentHTML", "onclick=", "onerror=",
                       "onload="):
            self.assertNotIn(banned, src,
                             f"UI 页面含被禁止的渲染通道：{banned}")
        # 渲染必须走 DOM API
        for required in ("createElement", "textContent",
                         "addEventListener"):
            self.assertIn(required, src)


# ---------------------------------------------------------------- 浏览器导入修复

class BrowserImportFixTests(unittest.TestCase):
    """/tmp/biaojing_p2_review 后续：浏览器导入修复任务书逐条回归。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="biaojing_p3_fix_")
        self.wb = ws_mod.Workbench(self.tmp)

    def tearDown(self):
        self.wb.close()

    def test_reject_record_api_and_length_caps(self):
        self.wb.record_rejected("x" * 600, "y" * 400)
        row = self.wb.conn.execute(
            "SELECT ref, reason FROM source_refs WHERE status='rejected'"
            " ORDER BY id DESC LIMIT 1").fetchone()
        self.assertEqual(len(row["ref"]), 512)
        self.assertEqual(len(row["reason"]), 300)

    def test_confirm_action_whitelist(self):
        # 未知 action 必须拒绝且不写库
        data = B01_DOCX.read_bytes()
        self.wb.ingest_bytes("蓝湾.docx", data)
        eid = self.wb.conn.execute(
            "SELECT evidence_id FROM candidates LIMIT 1").fetchone()["evidence_id"]
        before = self.wb.conn.execute(
            "SELECT COUNT(*) n FROM confirmations").fetchone()["n"]
        ok = self.wb.confirm_field("EV-1", "L1", "B1", "total_price",
                                   100.0, eid, "unexpected")
        self.assertFalse(ok["ok"])
        self.assertIn("未知操作", ok["error"])
        after = self.wb.conn.execute(
            "SELECT COUNT(*) n FROM confirmations").fetchone()["n"]
        self.assertEqual(before, after)  # 数据库无新行

    def test_confirm_locator_contract_enforced(self):
        # 存在但 locator 不符合 typed 契约的证据不得确认
        data = B01_DOCX.read_bytes()
        self.wb.ingest_bytes("蓝湾.docx", data)
        eid = self.wb.conn.execute(
            "SELECT evidence_id FROM candidates LIMIT 1").fetchone()["evidence_id"]
        # 人为把该证据的 locator_json 改成 kind=unknown（模拟坏库/坏数据）
        self.wb.conn.execute(
            "UPDATE evidence_store SET locator_json='{\"kind\": \"unknown\"}'"
            " WHERE evidence_id=?", (eid,))
        self.wb.conn.commit()
        ok = self.wb.confirm_field("EV-1", "L1", "B1", "total_price",
                                   100.0, eid, "confirm")
        self.assertFalse(ok["ok"])
        self.assertIn("定位无效", ok["error"])

    def test_total_price_rejects_bool_and_nonfinite(self):
        data = B01_DOCX.read_bytes()
        self.wb.ingest_bytes("蓝湾.docx", data)
        eid = self.wb.conn.execute(
            "SELECT evidence_id FROM candidates WHERE field='total_price'"
            " LIMIT 1").fetchone()["evidence_id"]
        for bad in (True, False, float("inf"), "abc"):
            ok = self.wb.confirm_field("EV-1", "L1", "B1", "total_price",
                                       bad, eid, "confirm")
            self.assertFalse(ok["ok"], repr(bad))
        rows = self.wb.conn.execute(
            "SELECT COUNT(*) n FROM confirmations").fetchone()["n"]
        self.assertEqual(rows, 0)
        # 数值、币种、税口径必须一次提交，失败时不留部分确认。
        for bad in (True, False, float("inf"), "abc"):
            ok = self.wb.confirm_amount("EV-1", "L1", "B1", bad, "yuan",
                                        "CNY", True, eid)
            self.assertFalse(ok["ok"], repr(bad))
        self.assertEqual(self.wb.conn.execute(
            "SELECT COUNT(*) FROM confirmations").fetchone()[0], 0)
        ok = self.wb.confirm_amount("EV-1", "L1", "B1", "100.5", "yuan",
                                    "CNY", True, eid)
        self.assertTrue(ok["ok"])
        vals = {}
        for row in self.wb.conn.execute("SELECT field,value FROM confirmations"):
            try:
                vals[row["field"]] = json.loads(row["value"])
            except (TypeError, ValueError):
                vals[row["field"]] = row["value"]
        self.assertEqual(vals["total_price"], "100.5")
        self.assertEqual(vals["currency"], "CNY")
        self.assertIs(vals["tax_included"], True)

    def test_foreign_or_conflicting_amount_is_rejected_without_partial_write(self):
        self.wb.ingest_bytes("蓝湾.docx", B01_DOCX.read_bytes())
        eid = self.wb.conn.execute(
            "SELECT evidence_id FROM candidates WHERE field='total_price'"
            " LIMIT 1").fetchone()["evidence_id"]
        for raw, currency in (("USD 100", "USD"), ("CNY $100", "CNY")):
            before = self.wb.conn.execute(
                "SELECT COUNT(*) FROM confirmations").fetchone()[0]
            result = self.wb.confirm_amount(
                "EV-1", "L1", "B1", raw, "yuan", currency, True, eid)
            self.assertFalse(result["ok"])
            self.assertEqual(self.wb.conn.execute(
                "SELECT COUNT(*) FROM confirmations").fetchone()[0], before)

    def test_build_facts_bool_price_not_number(self):
        # 规则事实边界直接拒绝布尔金额，不能等价于数字 1 再参与 R005。
        from biaojing.rules import FactBase, screen_all
        data = facts(
            event("EV-1", [lot("L1", [
                bid("B1", total_price=True, total_price_evidence="E-p1"),
                bid("B2", total_price=True, total_price_evidence="E-p2")])]),
            evidence=[ev("E-p1"), ev("E-p2")])
        with self.assertRaisesRegex(ValueError, "总报价"):
            screen_all(FactBase.from_dict(data))

    def test_run_screen_does_not_hide_invalid_confirmed_types(self):
        # 模拟旧库/坏库中的确认值：run_screen 必须把原类型交给 P2 校验，
        # 不得把布尔金额或字符串税口径悄悄改成 unknown。
        import datetime

        cases = (("total_price", "true", "总报价"),
                 ("tax_included", '"false"', "tax_included"))
        for field, raw_value, message in cases:
            with self.subTest(field=field):
                self.wb.conn.execute(
                    "INSERT INTO confirmations(event_id, lot_id, bidder_id,"
                    " field, value, action, confirmed_at) VALUES(?,?,?,?,?,?,?)",
                    ("EV-BAD", "L1", "B1", field, raw_value, "confirm",
                     datetime.datetime.now().isoformat()))
                self.wb.conn.commit()
                with self.assertRaisesRegex(ValueError, message):
                    self.wb.run_screen()
                self.wb.conn.execute(
                    "DELETE FROM confirmations WHERE event_id='EV-BAD'")
                self.wb.conn.commit()

    def test_evidence_refs_and_download_security(self):
        data = B01_DOCX.read_bytes()
        r = self.wb.ingest_bytes("蓝湾/含%2F百分号.docx", data)
        sha = r["sha256"]
        # 文件名单次解码：字面 %2F 保留在来源引用中，不被改写为路径分隔符
        self.assertIn("含%2F百分号", self.wb.refs_for_sha(sha)[0])
        # 原件字节按哈希路径完整保留
        self.assertEqual(self.wb.original_bytes(sha), data)
        expected_ref = f"蓝湾/含%2F百分号.docx（{len(data)} 字节）"
        self.assertIn(expected_ref, self.wb.refs_for_sha(sha))
        self.assertEqual(self.wb.display_name_for(sha), "含%2F百分号.docx")
        self.assertIsNone(self.wb.original_bytes("f" * 64))  # 不存在 → None

    def test_ui_walk_pagination_and_no_inline_handlers(self):
        from biaojing import ui_page

        src = ui_page.INDEX_HTML
        # readEntries 分页循环（大目录 101+ 条不截断）：空批判断必须存在
        self.assertIn("readEntries", src)
        self.assertIn("!batch.length", src)
        # 枚举失败上报端点被页面调用
        self.assertIn("/api/reject_record", src)
        # 空目录明确提示，不显示导入成功
        self.assertIn("未发现文件", src)


# ---------------------------------------------------------------- UI 行为回归（Node stub 执行）

class UploadBehaviorTests(unittest.TestCase):
    """真实浏览器复核（Chrome 点击选择 8 文件 → /api/state=0）的根因回归：
    1) drop 点击打开文件选择器后，input.click() 的合成 click 冒泡回 #drop
       重入 handler（两次 user-activation 警告）——target 守卫必须阻断；
    2) uploadFiles 对畸形项（如 undefined）单条可见失败，不中止批次。
    用 Node + DOM stub **实际执行**页面 JS，非静态字符串检查。"""

    @classmethod
    def setUpClass(cls):
        import shutil
        import subprocess
        if not shutil.which("node"):
            raise unittest.SkipTest("node 不可用")
        from biaojing.ui_page import INDEX_HTML
        m = __import__("re").search(
            r"<script>\n(.*?)</script>", INDEX_HTML, __import__("re").S)
        js = m.group(1)
        # /api/state 返回完整合法 state：页面末尾的 loadState() 正常完成，
        # 不产生 unhandled rejection，Node 进程生命周期完全确定
        state_stub = json.dumps({
            "product_name": "标镜", "findings": [{
                "rule_id": "R002", "rule_version": "P2-1.1",
                "scope": {"contact_kind": "phone",
                          "contact_value": "138-0000-0001",
                          "bidder_ids": ["SYN-BID-BLUE", "SYN-BID-RED"]},
                "inputs": {}, "params": {},
                "trigger_reason": "phone 138-0000-0001 出现在 2 个不同投标主体",
                "evidence_ids": ["E-aaa11112222", "E-bbb33334444"],
                "limitations": ["人工核实"], "alternative_explanations": ["代填"],
                "review_status": "pending_manual", "signal": "线索"}],
            "coverage": {
                "total_input_occurrences": 0, "unique_files": 0,
                "by_status": {},
                "status_order": ["success", "partial", "failed",
                                 "pending_ocr", "pending_convert",
                                 "duplicate", "rejected", "archived",
                                 "unknown"]},
            "source_rows": [{
                "ref": "160页投标文件.pdf", "doc_type": "pdf",
                "status": "partial", "reason": None, "parser": "pymupdf",
                "sha256": "a" * 64, "extract_version": "pymupdf test profile",
                "counts_json": json.dumps({"pages_ocr": 151,
                                            "pages_pending_ocr": 9})}],
            "source_total": 2, "source_offset": 0, "source_limit": 100,
            "source_has_more": True,
            "candidates": [],
            "candidate_total": 0, "candidate_offset": 0,
            "candidate_limit": 200, "candidate_has_more": False,
            "confirmation_count": 0, "confirmation_history": [],
            "confirmation_history_total": 0,
            "file_binding_history": [], "file_binding_history_total": 0,
            "last_screen_run": {"status": "completed",
                                "rule_statuses": {}}})
        harness = """
const src=%s;
const STATE=%s;
const EV={evidence_id:"E-deadbeefcafe-0008",sha256:"d".repeat(64),
  locator_display:"paragraph 8 image 1",source_type:"unknown",
  page_status:"ocr",page_note:"需对照原始 Word 图片复核",
  ocr_used:true,
  refs:["蓝湾.docx（123 字节）"],
  quote:"【合成】证据首句……（完整不截断）……证据末句"};
const calls=[]; const clickLog=[];
const els={};
function makeEl(){const n={_kids:[],_handlers:{},textContent:"",style:{},
  dataset:{},value:"",
  appendChild(c){n._kids.push(c);return c;},
  insertBefore(c){n._kids.unshift(c);return c;},
  addEventListener(t,f){n._handlers[t]=f;},remove(){},
  classList:{add(){},remove(){}},click(){clickLog.push("click");}};
  return n;}
global.document={querySelector:s=>(els[s]=els[s]||makeEl()),
  createElement:()=>makeEl(),
  createTextNode:t=>({textContent:t,_kids:[]})};
global.window={open(){}};
global.setTimeout=()=>0;
global.fetch=async(url,opt)=>{calls.push({url,opt:opt||{}});
  if(url.indexOf("/api/state")===0){
    const state=JSON.parse(JSON.stringify(STATE));
    const sourceOffset=Number(new URL(url,"http://local").searchParams.get("source_offset")||0);
    if(sourceOffset){state.source_rows=[{ref:"Word投标文件.docx",doc_type:"docx",
      status:"partial",parser:"python-docx",sha256:"b".repeat(64),
      extract_version:"docx image OCR",counts_json:JSON.stringify({images_ocr:2,
      images_pending_ocr:1})}];
      state.source_offset=sourceOffset;state.source_has_more=false;}
    return {ok:true,status:200,json:async()=>state};
  }
  if(url.indexOf("/api/evidence/")===0)
    return {ok:true,status:200,json:async()=>EV};
  if(url==="/api/screen")
    return {ok:false,status:400,json:async()=>({error:"确认数据无法筛查：类型无效"})};
  if(url==="/api/confirm")
    return {ok:true,status:200,json:async()=>({ok:true})};
  if(url==="/api/retry_ocr")
    return {ok:true,status:200,json:async()=>({job_id:"job-test"})};
  if(url==="/api/jobs/job-test")
    return {ok:true,status:200,json:async()=>({job_id:"job-test",
      status:"completed",page:151,total_pages:151,stage:"已完成",
      result:{retried_pages:9,pages_pending_ocr:0}})};
  if(url==="/api/jobs/job-docx")
    return {ok:true,status:200,json:async()=>({job_id:"job-docx",
      kind:"docx_upload",status:"completed",page:1,total_pages:2,
      stage:"page_done",result:{status:"success"}})};
  return {ok:true,status:200,
    json:async()=>({result:{status:"success",ref:"x"}})}};
// 真实分隔符用 String.fromCharCode(10) 生成，避免 Python→JS 的多层
// 反斜杠转义；eval 内同时覆盖 loadState，消除与页面自身调用的竞态
const NL=String.fromCharCode(10);
eval(src+NL+";globalThis.__T={uploadFiles,reportReject,dropClickHandler,els,calls,clickLog,renderCands,loadState,setLoadState:fn=>{loadState=fn}};"
  + NL+"globalThis.__T.waitForJob=waitForJob;"
  + NL+"loadState=async()=>{};");
""" % (json.dumps(js), state_stub)
        harness += """
function collect(n,out){out=out||[];if(!n||!n._kids)return out;
  out.push(n);for(const k of n._kids)collect(k,out);return out}
(async()=>{
  const T=globalThis.__T;
  const good={f:{name:"a.txt",size:10},rel:"a.txt"};
  await T.uploadFiles([undefined,good]);
  const dropEl=T.els["#drop"], fileEl=T.els["#file"];
  T.dropClickHandler({target:dropEl});          // 点击落在拖放区本身 → 应打开
  T.dropClickHandler({target:fileEl});          // 合成 click 冒泡 → 不得重开
  T.dropClickHandler({target:fileEl});
  // 候选行「查看证据」：确认前即可打开完整证据面板
  const state={product_name:"标镜",
    coverage:{total_input_occurrences:0,unique_files:0,by_status:{},
      status_order:["success"]},
    source_rows:[],source_total:0,source_has_more:false,
    candidates:[{id:8,sha256:"d".repeat(64),field:"total_price",
      value:119860000,evidence_id:"E-deadbeefcafe-0008",
      locator_display:"paragraph 8",note:"标签匹配",companion:null}],
    confirmation_count:0};
  T.renderCands(state);
  const btns=collect(T.els["#cands"]).filter(
    n=>n.textContent==="查看证据"&&n._handlers&&n._handlers.click);
  if(btns.length!==1)
    throw new Error("查看证据按钮数量异常:"+btns.length);
  await btns[0]._handlers.click();
  const panel=T.els["#findings"]._kids[0];
  const panelText=collect(panel).map(n=>n.textContent||"").join("│");
  // 刷新回显：执行页面原版 loadState（/api/state stub 含非空 findings）
  await T.loadState();
  if(T.els["#moreSources"].hidden)
    throw new Error("来源分页按钮未显示");
  T.setLoadState(T.loadState);
  await T.els["#moreSources"]._handlers.click();
  if(!T.els["#sourcePage"].textContent.includes("2 / 2"))
    throw new Error("来源追加分页未保留并显示全部来源");
  const fbKids=collect(T.els["#findings"]);
  const findingsBox=fbKids.map(n=>n.textContent||"").join("│");
  const sourceText=collect(T.els["#sources"]).map(n=>n.textContent||"").join("│");
  const retryBtns=collect(T.els["#sources"]).filter(
    n=>/^重试待 OCR/.test(n.textContent||"")&&n._handlers&&n._handlers.click);
  if(!retryBtns.length)throw new Error("待 OCR PDF 未显示重试按钮");
  const retryButtonText=retryBtns[0].textContent;
  await retryBtns[0]._handlers.click();
  const retryCalls=calls.filter(c=>c.url==="/api/retry_ocr");
  const jobPollCalls=calls.filter(c=>c.url==="/api/jobs/job-test");
  await T.els["#runScreen"]._handlers.click();
  const screenError=T.els["#toast"].textContent;
  await T.waitForJob("job-docx","Word OCR");
  const docxTaskText=T.els["#taskText"].textContent;
  const confirmCallsBeforePriceEditor=calls.filter(c=>c.url==="/api/confirm").length;
  for(const selector of ["#eid","#lid","#bid"])
    T.els[selector]=T.els[selector]||makeEl();
  T.els["#eid"].value="EV-SYN-001";T.els["#lid"].value="SYN-LOT-001";
  T.els["#bid"].value="SYN-BIDDER-01";
  T.els["#cands"]._kids=[];
  T.renderCands({candidates:[{id:42,field:"price_lines",value:[{
    item_code:"A1",item_name:"电线",spec:"BV",unit:"unknown",
    unit_price:null,unit_price_raw:"100",amount_unit:"unknown",
    currency:"unknown",tax_included:null,evidence:"E-line-42",
    field_evidence:{unit_price:"E-line-42"}}],
    evidence_id:"E-line-42",locator_display:"报价!H2",note:"金额单位待核"}],
    candidate_total:1,candidate_has_more:false});
  const priceControls=collect(T.els["#cands"]).filter(n=>n.dataset
    &&n.dataset.priceLineField);
  const controlByField=Object.fromEntries(priceControls.map(n=>
    [n.dataset.priceLineField,n]));
  controlByField.unit.value="台";controlByField.unit_price_raw.value="100";
  controlByField.amount_unit.value="ten_thousand_yuan";
  controlByField.currency.value="CNY";controlByField.tax_included.value="true";
  const correctButton=collect(T.els["#cands"]).find(n=>n.textContent==="更正"
    &&n._handlers&&n._handlers.click);
  await correctButton._handlers.click();
  const priceConfirm=calls.filter(c=>c.url==="/api/confirm").slice(-1)[0];
  const priceLineValue=JSON.parse(priceConfirm.opt.body).value[0];
  const advancedLine={item_code:"A2",item_name:"钢筋",spec:"HRB400",
    unit:"unknown",unit_price:null,unit_price_raw:"200",amount_unit:"unknown",
    currency:"unknown",tax_included:null,evidence:"E-line-43",
    field_evidence:{unit_price:"E-line-43"}};
  T.els["#cands"]._kids=[];
  T.renderCands({candidates:[{id:43,field:"price_lines",value:[advancedLine],
    evidence_id:"E-line-43",locator_display:"报价!H3",note:"金额单位待核"}],
    candidate_total:1,candidate_has_more:false});
  const advancedControls=collect(T.els["#cands"]).filter(n=>n.dataset
    &&n.dataset.priceLineField);
  const advancedByField=Object.fromEntries(advancedControls.map(n=>
    [n.dataset.priceLineField,n]));
  advancedByField.unit.value="控件单位";
  const advancedJson=collect(T.els["#cands"]).find(n=>n.dataset
    &&n.dataset.cid==43);
  advancedJson.value=JSON.stringify([{...advancedLine,unit:"JSON单位"}],null,2);
  const advancedButton=collect(T.els["#cands"]).find(n=>n.textContent==="更正"
    &&n._handlers&&n._handlers.click);
  await advancedButton._handlers.click();
  const advancedConfirm=calls.filter(c=>c.url==="/api/confirm").slice(-1)[0];
  const advancedLineValue=JSON.parse(advancedConfirm.opt.body).value[0];
  console.log(JSON.stringify({calls:T.calls,
    log:T.els["#uploadLog"].textContent, clickLog:T.clickLog,
    evCalls:calls.filter(c=>c.url.indexOf("/api/evidence/")===0)
      .map(c=>c.url),
    confirmCalls:confirmCallsBeforePriceEditor,priceLineValue,advancedLineValue,
    panelText, panelCount:T.els["#findings"]._kids.length,
    findingsBox, findingsPanelCount:fbKids.length,sourceText,retryButtonText,
    retryCalls:retryCalls.length,jobPollCalls:jobPollCalls.length,screenError,
    docxTaskText}));
})().catch(e=>{console.error("HARNESS_FAIL",e&&e.stack||e&&e.message);process.exit(1)});
"""
        cls.tmp = tempfile.mkdtemp(prefix="biaojing_p3_behavior_")
        cls.js_path = os.path.join(cls.tmp, "page.js")
        Path(cls.js_path).write_text(harness, encoding="utf-8")
        cls.proc = subprocess.run(
            ["node", cls.js_path], capture_output=True, text=True, timeout=30)

    def test_harness_ran(self):
        self.assertNotIn("HARNESS_FAIL", self.proc.stderr,
                         self.proc.stderr[:400])

    def test_malformed_item_fails_visibly_batch_continues(self):
        payload = json.loads(self.proc.stdout)
        uploads = [c for c in payload["calls"]
                   if c["url"].startswith("/api/upload")]
        self.assertEqual(len(uploads), 1, "畸形项不得发起上传")
        self.assertIn("不是有效文件对象", payload["log"])
        self.assertIn("[成功]", payload["log"].replace("success", "成功")
                      ) or self.assertIn("success", payload["log"])
        self.assertIn("已处理 1，失败/拒收 1", payload["log"])
        self.assertIn("a.txt", payload["log"])

    def test_synthetic_click_bubble_does_not_reopen_chooser(self):
        payload = json.loads(self.proc.stdout)
        self.assertEqual(payload["clickLog"].count("click"), 1,
                         "合成 click 冒泡回 #drop 不得再次打开选择器")

    def test_refresh_reloads_findings_after_loadstate(self):
        # 刷新回显缺陷回归：loadState 必须把 /api/state 的 findings 数组
        # 以 {findings:[...]} 契约传给 renderFindings——刷新后 findings
        # 容器含 R002 与证据 ID，不再显示"暂无筛查结果"。
        # 在 Node 中执行页面**原版 loadState**（真实渲染路径）。
        payload = json.loads(self.proc.stdout)
        box = payload["findingsBox"]
        self.assertIn("R002", box)
        self.assertIn("E-aaa11112222", box)
        self.assertIn("E-bbb33334444", box)
        self.assertNotIn("暂无筛查结果", box)

    def test_source_row_shows_ocr_and_pending_page_counts(self):
        payload = json.loads(self.proc.stdout)
        self.assertIn("本机 OCR 成功 151 页，待 OCR 9 页",
                      payload["sourceText"])

    def test_pending_pdf_shows_retry_button_and_calls_retry_route(self):
        payload = json.loads(self.proc.stdout)
        self.assertEqual(payload["retryButtonText"], "重试待 OCR 9 页")
        self.assertEqual(payload["retryCalls"], 1)
        self.assertEqual(payload["jobPollCalls"], 1)

    def test_word_image_evidence_and_background_progress_are_labeled(self):
        payload = json.loads(self.proc.stdout)
        self.assertIn("paragraph 8 image 1", payload["panelText"])
        self.assertIn("页面状态：│ocr", payload["panelText"])
        self.assertIn("需对照原始 Word 图片复核", payload["panelText"])
        self.assertIn("1 / 2 张图片", payload["docxTaskText"])
        self.assertIn("本机 OCR 成功 2 张图片，待 OCR 1 张图片",
                      payload["sourceText"])

    def test_screen_error_is_visible_without_js_exception(self):
        payload = json.loads(self.proc.stdout)
        self.assertIn("确认数据无法筛查：类型无效", payload["screenError"])

    def test_upload_log_contains_real_newlines(self):
        # 独立复核：上传日志必须是真实换行（textContent 含 chr(10)），
        # 且每个行边界都不得以字面 backslash+n（chr(92)+chr(110)，过度
        # 转义会使日志显示 "\n" 字面量）结尾。本断言在 Node 中实际执行
        # 页面 JS 后检查运行时值；转义全部用 chr() 构造，避免本测试自身
        # 再引入反斜杠转义歧义。
        payload = json.loads(self.proc.stdout)
        log = payload["log"]
        NL, BS, N = chr(10), chr(92), chr(110)
        self.assertIn(NL, log, "日志应为真实换行分隔")
        self.assertNotIn(BS + N, log, "日志含字面 backslash+n（过度转义）")
        for i, line in enumerate(log.split(NL), start=1):
            self.assertFalse(
                line.endswith(BS + N),
                f"第 {i} 行边界以字面 backslash+n 结尾：{line[-40:]!r}")

    def test_candidate_row_view_evidence_before_confirm(self):
        # UI 验收缺口回归：候选行「查看证据」在**确认之前**即可 fetch 并
        # 显示完整证据面板（原文全文/SHA-256/来源引用/下载按钮），
        # 且不产生任何 /api/confirm 调用
        payload = json.loads(self.proc.stdout)
        self.assertEqual(payload["evCalls"],
                         ["/api/evidence/E-deadbeefcafe-0008"])
        self.assertEqual(payload["confirmCalls"], 0)
        self.assertGreaterEqual(payload["panelCount"], 1)
        pt = payload["panelText"]
        self.assertIn("证据详情 E-deadbeefcafe-0008", pt)
        self.assertIn("【合成】证据首句", pt)
        self.assertIn("证据末句", pt)          # 完整原文不截断
        self.assertIn("d" * 64, pt)            # 文件 SHA-256
        self.assertIn("蓝湾.docx（123 字节）", pt)  # 来源引用
        self.assertIn("下载原始文件", pt)
        self.assertIn("paragraph 8", pt)       # 定位
        self.assertIn("本机 OCR 机器识别文本，请对照原件复核", pt)

    def test_price_line_unit_editor_submits_structured_values(self):
        # P2-10：逐行金额单位/计量单位/币种/税口径通过控件写回完整行，
        # 不要求用户手改整段 JSON。
        line = json.loads(self.proc.stdout)["priceLineValue"]
        self.assertEqual(line["unit"], "台")
        self.assertEqual(line["unit_price_raw"], "100")
        self.assertEqual(line["amount_unit"], "ten_thousand_yuan")
        self.assertEqual(line["currency"], "CNY")
        self.assertIs(line["tax_included"], True)

    def test_advanced_price_json_is_not_overwritten_by_stale_controls(self):
        # 高级 JSON 修改整行后，初始结构化控件不能再覆盖它。
        line = json.loads(self.proc.stdout)["advancedLineValue"]
        self.assertEqual(line["unit"], "JSON单位")

    def test_page_served_with_no_store(self):
        from biaojing import webapp
        import threading
        srv = webapp.WorkbenchServer(
            ("127.0.0.1", 0), ws_mod.Workbench(
                tempfile.mkdtemp(prefix="biaojing_p3_ns_")))
        port = srv.server_address[1]
        t = threading.Thread(target=srv.serve_forever, daemon=True)
        t.start()
        try:
            with urllib.request.urlopen(f"http://127.0.0.1:{port}/") as r:
                self.assertEqual(r.headers.get("Cache-Control"), "no-store")
        finally:
            srv.shutdown()
            srv.server_close()
            srv.workbench.close()


if __name__ == "__main__":
    unittest.main()
