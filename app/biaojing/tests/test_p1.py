# -*- coding: utf-8 -*-
"""P1 底座测试：全部使用临时目录内合成的 PDF/DOCX/XLSX，不触碰真实业务文件。

运行：
    cd app && python -m unittest biaojing.tests.test_p1 -v
"""

from __future__ import annotations

import base64
import io
import json
import os
from pathlib import Path
import shutil
import struct
import tempfile
import unittest
import zipfile
from unittest.mock import patch

import pymupdf
import docx as docx_lib
import openpyxl

from biaojing import cli as biaojing_cli
from biaojing import hashing, xlsx_parser
from biaojing.walker import iter_inputs
from biaojing.zip_container import expand, guard_ooxml

# 1x1 红色 PNG（合成图像素材，用于扫描页与 DOCX 内嵌图）
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108020000009077"
    "53de0000000c4944415408d763f8cfc0000000030001"
    "7dd2b7450000000049454e44ae426082"
)


def _tmp(name: str) -> str:
    return os.path.join(tempfile.mkdtemp(prefix="biaojing_t_"), name)


def make_pdf(path: str, pages: list[str], encrypt: bool = False) -> None:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text, fontsize=12, fontname="china-s")
    if encrypt:
        doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_128,
                 owner_pw="o", user_pw="u")
    else:
        doc.save(path)
    doc.close()


def make_scan_pdf(path: str, pages: list[str], image_pages: set[int]) -> None:
    """合成 PDF：指定页插入整页图像且无文本（扫描页），其余页纯文本。"""
    doc = pymupdf.open()
    for i, text in enumerate(pages, start=1):
        page = doc.new_page()
        if i in image_pages:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 60), 0)
            page.insert_image(page.rect, pixmap=pix)
        if text:
            page.insert_text((72, 72), text, fontsize=12)
    doc.save(path)
    doc.close()


def make_docx(path: str, paragraphs: list[str],
              tables: list[list[list[str]]] | None = None,
              header: str | None = None,
              inline_png: bool = False) -> None:
    doc = docx_lib.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    for table in tables or []:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    if header is not None:
        doc.sections[0].header.paragraphs[0].text = header
    if inline_png:
        doc.add_picture(io.BytesIO(PNG_1PX))
    doc.save(path)


def make_xlsx(path: str, sheets: dict[str, list[tuple]],
              hidden: dict[str, str] | None = None,
              comment: tuple[str, str, str] | None = None) -> None:
    """sheets: {sheet名: [(坐标, 值), ...]}；hidden: {sheet名: state}；
    comment: (sheet名, 坐标, 批注文本)"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(title=name)
        for coord, val in cells:
            ws[coord] = val
    for name, state in (hidden or {}).items():
        wb[name].sheet_state = state
    if comment:
        from openpyxl.comments import Comment

        s, coord, text = comment
        wb[s][coord].comment = Comment(text, "tester")
    wb.save(path)


def inject_cached_value(path: str, sheet_xml: str, cell_ref: str, value: str) -> None:
    """把 xlsx 内某公式单元格的 XML 注入缓存值 <v>，模拟带缓存的真实文件。"""
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_xml:
                text = data.decode("utf-8")
                old = f'<c r="{cell_ref}"><f>'
                new = f'<c r="{cell_ref}"><f>'
                assert old in text, f"未找到公式单元格 {cell_ref}"
                text = text.replace(old, new, 1)
                # 在该 </f> 后插入缓存值
                idx = text.find(f'<c r="{cell_ref}"><f>')
                end = text.find("</f>", idx) + len("</f>")
                text = text[:end] + f"<v>{value}</v>" + text[end:]
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, path)


def run_cli(*paths, **opts) -> dict:
    argv = [str(p) for p in paths]
    out = opts.pop("output", None)
    if out:
        argv += ["-o", str(out)]
    for k, v in opts.items():
        argv += [f"--{k.replace('_', '-')}", str(v)]
    rc = biaojing_cli.run(argv)
    assert rc == 0
    with open(out, encoding="utf-8") as f:
        return json.load(f)


class Sha256Tests(unittest.TestCase):
    def test_file_hash_matches_independent_hashlib(self):
        path = _tmp("a.txt")
        payload = "标镜测试字节".encode("utf-8") + os.urandom(64)
        with open(path, "wb") as f:
            f.write(payload)
        import hashlib
        self.assertEqual(hashing.sha256_file(path),
                         hashlib.sha256(payload).hexdigest())


class PdfTests(unittest.TestCase):
    def test_per_page_text_and_real_page_numbers(self):
        path = _tmp("multi.pdf")
        make_pdf(path, ["第一页甲", "第二页乙", "第三页丙"])
        from biaojing import pdf_parser

        result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["counts"]["pages_total"], 3)
        self.assertEqual([e["page"] for e in result["evidence"]], [1, 2, 3])
        self.assertIn("第二页乙", result["evidence"][1]["text"])
        self.assertEqual(result["evidence"][1]["locator"], "page 2")

    def test_scan_page_pending_ocr_and_mixed_partial(self):
        path = _tmp("scan.pdf")
        make_scan_pdf(path, ["", "正常文本页", ""], image_pages={1, 3})
        from biaojing import pdf_parser

        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "_ocr_page", return_value=("", None)):
            result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["counts"]["pages_pending_ocr"], 2)
        self.assertEqual(result["counts"]["pages_text"], 1)
        self.assertEqual(result["evidence"][0]["page_status"], "pending_ocr")
        self.assertEqual(result["evidence"][0]["page"], 1)

    def test_mixed_text_and_scanned_image_page_runs_ocr(self):
        doc = pymupdf.open()
        page = doc.new_page(width=360, height=180)
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 60), 0)
        page.insert_image(page.rect, pixmap=pix)
        page.insert_text((20, 28), "表头", fontsize=12, fontname="china-s")
        data = doc.tobytes()
        doc.close()
        from biaojing import pdf_parser

        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "_ocr_page",
                             return_value=("扫描内容", None)) as ocr:
            result = pdf_parser.parse(data)
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual(result["evidence"][0]["page_status"], "ocr")
        self.assertIn("表头", result["evidence"][0]["text"])
        self.assertIn("扫描内容", result["evidence"][0]["text"])
        self.assertIn("混合页面", result["evidence"][0]["page_note"])
        self.assertEqual(result["counts"]["pages_pending_ocr"], 0)

        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=(None, None, "缺少 Tesseract")):
            pending = pdf_parser.parse(data)
        self.assertEqual(pending["evidence"][0]["page_status"], "pending_ocr")
        self.assertEqual(pending["evidence"][0]["text"], "表头")
        self.assertIn("混合页面", pending["evidence"][0]["page_error"])

    def test_all_scan_pages_pending_ocr(self):
        path = _tmp("allscan.pdf")
        make_scan_pdf(path, ["", ""], image_pages={1, 2})
        from biaojing import pdf_parser

        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "_ocr_page", return_value=("", None)):
            result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "pending_ocr")
        self.assertEqual(result["counts"]["pages_text"], 0)

    def test_scanned_pdf_gets_local_ocr_text(self):
        from biaojing import pdf_parser

        binary, lang, _ = pdf_parser._ocr_runtime()
        if not binary or not lang:
            self.skipTest("本机缺少 Tesseract 或 chi_sim 中文模型")
        source = pymupdf.open()
        p = source.new_page(width=360, height=180)
        p.insert_text((32, 92), "OCR TEST 12345", fontsize=28,
                      fontname="helv")
        pix = p.get_pixmap(matrix=pymupdf.Matrix(200 / 72, 200 / 72),
                           alpha=False)
        scanned = pymupdf.open()
        scan_page = scanned.new_page(width=360, height=180)
        scan_page.insert_image(scan_page.rect, pixmap=pix)
        data = scanned.tobytes()
        source.close()
        scanned.close()

        result = pdf_parser.parse(data)
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["counts"]["pages_ocr"], 1)
        self.assertEqual(result["counts"]["pages_pending_ocr"], 0)
        evidence = result["evidence"][0]
        self.assertEqual(evidence["page_status"], "ocr")
        self.assertTrue(evidence["ocr_used"])
        self.assertIn("OCR", evidence["text"])
        self.assertIn("12345", evidence["text"])

    def test_embedded_image_ocr_checks_dimensions_before_decode(self):
        from biaojing import pdf_parser

        oversized = (b"\x89PNG\r\n\x1a\n" + b"\x00\x00\x00\x0dIHDR" +
                     (50_000).to_bytes(4, "big") + (1_000).to_bytes(4, "big"))
        text, error = pdf_parser.ocr_image(oversized, "unused", "chi_sim")
        self.assertEqual(text, "")
        self.assertIn("像素超出", error)
        text, error = pdf_parser.ocr_image(b"unsupported", "unused", "chi_sim")
        self.assertEqual(text, "")
        self.assertIn("无法安全预检", error)

    def test_embedded_png_jpeg_gif_bmp_decode_to_ocr_input(self):
        from biaojing import pdf_parser

        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 30, 20), 0)
        gif = base64.b64decode("R0lGODlhAQABAIAAAAAAAP///ywAAAAAAQABAAACAUwAOw==")
        bmp = (b"BM" + struct.pack("<IHHI", 58, 0, 0, 54) +
               struct.pack("<IiiHHIIiiII", 40, 1, 1, 1, 24, 0, 4, 0, 0, 0, 0) +
               bytes(4))
        images = (pix.tobytes("png"), pix.tobytes("jpeg"), gif, bmp)
        with patch.object(pdf_parser, "_ocr_png", return_value=("decoded", None)) as run:
            for image in images:
                self.assertEqual(pdf_parser.ocr_image(image, "tesseract", "eng"),
                                 ("decoded", None))
        self.assertEqual(run.call_count, 4)

    def test_word_embedded_image_real_local_ocr(self):
        from biaojing import docx_parser, pdf_parser

        binary, lang, _ = pdf_parser._ocr_runtime()
        if not binary or not lang:
            self.skipTest("本机缺少 Tesseract 或 chi_sim 中文模型")
        source = pymupdf.open()
        page = source.new_page(width=900, height=240)
        page.insert_text((24, 90), "投标人名称：华夏建设工程有限公司",
                         fontname="china-s", fontsize=30)
        page.insert_text((24, 160), "项目编号：IMAGE OCR SAMPLE 2026",
                         fontname="china-s", fontsize=27)
        image = page.get_pixmap(matrix=pymupdf.Matrix(2, 2),
                                alpha=False).tobytes("png")
        source.close()
        path = _tmp("word_ocr.docx")
        doc = docx_lib.Document()
        doc.add_paragraph().add_run().add_picture(io.BytesIO(image))
        doc.save(path)

        result = docx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "success")
        evidence = next(e for e in result["evidence"] if e["kind"] == "docx_image")
        self.assertEqual(evidence["page_status"], "ocr")
        self.assertTrue(evidence["ocr_used"])
        self.assertIn("投标人名称", evidence["text"])
        self.assertIn("华夏建设工程有限公司", evidence["text"])

    def test_ocr_scans_more_than_100_pages(self):
        from biaojing import pdf_parser

        path = _tmp("ocr-over-100.pdf")
        page_count = 101
        make_scan_pdf(path, [""] * page_count,
                      image_pages=set(range(1, page_count + 1)))
        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "_ocr_page",
                             return_value=("机器识别文本", None)):
            result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["counts"]["pages_total"], page_count)
        self.assertEqual(result["counts"]["pages_ocr"], page_count)
        self.assertEqual(result["counts"]["pages_pending_ocr"], 0)
        self.assertEqual(len(result["evidence"]), page_count)
        self.assertTrue(all(item["ocr_used"] for item in result["evidence"]))

    def test_ocr_unavailable_and_pixel_cap_stay_explicit(self):
        from biaojing import pdf_parser

        path = _tmp("ocr-unavailable.pdf")
        make_scan_pdf(path, [""], image_pages={1})
        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=(None, None, "本机未安装 Tesseract，扫描页仍待 OCR")):
            missing = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(missing["status"], "pending_ocr")
        self.assertTrue(any("未安装 Tesseract" in n for n in missing["notes"]))
        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "OCR_MAX_PIXELS", 1):
            oversized = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(oversized["status"], "pending_ocr")
        self.assertTrue(any("像素超出 OCR 上限" in n
                            for n in oversized["notes"]))

    def test_blank_page_detected(self):
        path = _tmp("blank.pdf")
        make_pdf(path, ["有字", ""])
        from biaojing import pdf_parser

        result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["counts"]["pages_blank"], 1)

    def test_encrypted_pdf_fails(self):
        path = _tmp("enc.pdf")
        make_pdf(path, ["秘密"], encrypt=True)
        from biaojing import pdf_parser

        result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "failed")
        self.assertIn("加密", result["error"])

    def test_corrupt_pdf_fails_without_raising(self):
        path = _tmp("corrupt.pdf")
        with open(path, "wb") as f:
            f.write(b"%PDF-1.7 garbage not a pdf \x00\x01")
        from biaojing import pdf_parser

        result = pdf_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "failed")


class DocxTests(unittest.TestCase):
    def test_paragraph_and_table_numbering(self):
        path = _tmp("a.docx")
        make_docx(path,
                  paragraphs=["段一", "", "段三"],
                  tables=[[["R1C1", "R1C2"], ["R2C1", "R2C2"]]])
        from biaojing import docx_parser

        result = docx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "success")
        paras = [e for e in result["evidence"] if e["kind"] == "docx_paragraph"]
        self.assertEqual([p["paragraph_index"] for p in paras], [1, 3])
        cells = [e for e in result["evidence"] if e["kind"] == "docx_table_cell"]
        self.assertEqual(
            [(c["table_index"], c["row"], c["col"]) for c in cells],
            [(1, 1, 1), (1, 1, 2), (1, 2, 1), (1, 2, 2)])
        self.assertEqual(cells[0]["locator"], "table 1!R1C1")
        self.assertEqual(result["counts"]["paragraphs_total"], 3)
        self.assertEqual(result["counts"]["paragraphs_nonempty"], 2)

    def test_header_footer_extracted(self):
        path = _tmp("h.docx")
        make_docx(path, ["正文"], header="页眉标注")
        from biaojing import docx_parser

        result = docx_parser.parse(Path(path).read_bytes())
        hf = [e for e in result["evidence"] if e["kind"] == "docx_header_paragraph"]
        self.assertTrue(hf and hf[0]["text"] == "页眉标注")
        self.assertEqual(hf[0]["locator"], "section 1 header paragraph 1")

    def test_image_only_docx_pending_ocr(self):
        path = _tmp("img.docx")
        make_docx(path, paragraphs=[], inline_png=True)
        from biaojing import docx_parser

        result = docx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "pending_ocr")

    def test_docx_image_ocr_evidence_candidates_and_locations(self):
        path = _tmp("ocr_images.docx")
        doc = docx_lib.Document()
        doc.add_paragraph("投标文件")
        para = doc.add_paragraph()
        para.add_run().add_picture(io.BytesIO(PNG_1PX))
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(PNG_1PX))
        doc.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(PNG_1PX))
        doc.sections[0].footer.paragraphs[0].add_run().add_picture(io.BytesIO(PNG_1PX))
        doc.save(path)
        from biaojing import candidates, docx_parser, pdf_parser

        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "ocr_image",
                             return_value=("投标人名称：华夏建设有限公司", None)) as ocr:
            result = docx_parser.parse(Path(path).read_bytes())
        images = [e for e in result["evidence"] if e["kind"] == "docx_image"]
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["counts"]["images_ocr"], 4)
        self.assertEqual([e["locator"] for e in images], [
            "paragraph 2 image 1", "table 1!R1C1 image 1",
            "section 1 header paragraph 1 image 1",
            "section 1 footer paragraph 1 image 1"])
        self.assertEqual([e["image_index"] for e in images], [1, 2, 3, 4])
        self.assertEqual(ocr.call_count, 3)  # 同一媒体部件复用 OCR 结果
        tagged = candidates.assign_evidence_ids("ab" * 32, result["evidence"])
        bidder = [c for c in candidates.extract_candidates("ab" * 32, "docx", tagged)
                  if c["field"] == "bidder_name"]
        self.assertEqual(len(bidder), 4)
        self.assertEqual(bidder[0]["locator"]["kind"], "docx_image")
        self.assertEqual(bidder[0]["locator_display"], "paragraph 2 image 1")
        from biaojing.rules import _locator_valid
        self.assertTrue(_locator_valid(bidder[0]["locator"]))
        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "ocr_image",
                             return_value=("", "模拟 OCR 失败")):
            partial = docx_parser.parse(Path(path).read_bytes())
        self.assertEqual(partial["status"], "partial")
        self.assertEqual(partial["counts"]["images_pending_ocr"], 4)

    def test_docx_image_cancel_leaves_unprocessed_image_pending(self):
        path = _tmp("cancel_images.docx")
        doc = docx_lib.Document()
        doc.add_paragraph().add_run().add_picture(io.BytesIO(PNG_1PX))
        other_png = pymupdf.Pixmap(
            pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), 0).tobytes("png")
        doc.add_paragraph().add_run().add_picture(io.BytesIO(other_png))
        doc.save(path)
        from biaojing import docx_parser, pdf_parser

        completed = []
        def progress(index, total, stage, status):
            if stage == "page_done":
                completed.append(index)
        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "ocr_image",
                             return_value=("有足够长度的模拟 OCR 结果", None)) as ocr:
            result = docx_parser.parse(Path(path).read_bytes(),
                                      progress=progress,
                                      cancelled=lambda: len(completed) > 0)
        images = [e for e in result["evidence"] if e["kind"] == "docx_image"]
        self.assertTrue(result["cancelled"])
        self.assertEqual(ocr.call_count, 1)
        self.assertEqual([e["page_status"] for e in images], ["ocr", "pending_ocr"])
        self.assertIn("用户取消处理", images[1]["page_error"])

    def test_docx_image_total_byte_limit_stays_partial_with_reason(self):
        path = _tmp("limited_images.docx")
        doc = docx_lib.Document()
        doc.add_paragraph("可读取的正文")
        doc.add_paragraph().add_run().add_picture(io.BytesIO(PNG_1PX))
        doc.save(path)
        from biaojing import docx_parser, pdf_parser

        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(docx_parser, "OCR_MAX_TOTAL_IMAGE_BYTES", 1), \
                patch.object(pdf_parser, "ocr_image") as ocr:
            result = docx_parser.parse(Path(path).read_bytes())
        image = next(e for e in result["evidence"] if e["kind"] == "docx_image")
        self.assertEqual(result["status"], "partial")
        self.assertEqual(image["page_status"], "pending_ocr")
        self.assertIn("总量超出上限", image["page_error"])
        ocr.assert_not_called()

    def test_empty_docx_fails(self):
        path = _tmp("empty.docx")
        make_docx(path, [])
        from biaojing import docx_parser

        result = docx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "failed")


class XlsxTests(unittest.TestCase):
    def test_values_locator_format_comment(self):
        path = _tmp("a.xlsx")
        make_xlsx(path,
                  sheets={"报价": [("A1", 100.5), ("B2", "说明")]},
                  comment=("报价", "B2", "此处需要人工核对"))
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "success")
        by_locator = {e["locator"]: e for e in result["evidence"]
                      if e["kind"] == "xlsx_cell"}
        self.assertIn("报价!A1", by_locator)
        self.assertEqual(by_locator["报价!A1"]["value"], 100.5)
        self.assertEqual(by_locator["报价!B2"]["comment"], "此处需要人工核对")
        self.assertTrue(by_locator["报价!A1"]["number_format"])

    def test_formula_without_cached_value_is_unknown(self):
        path = _tmp("f.xlsx")
        make_xlsx(path, sheets={"S": [("A1", 1), ("A2", 2), ("A3", "=SUM(A1:A2)")]})
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes())
        cell = [e for e in result["evidence"] if e.get("cell") == "A3"][0]
        self.assertEqual(cell["formula"], "=SUM(A1:A2)")
        self.assertEqual(cell["cached_value"], "unknown")
        self.assertEqual(result["counts"]["formulas_cached_unknown"], 1)

    def test_formula_with_cached_value_read(self):
        path = _tmp("fc.xlsx")
        make_xlsx(path, sheets={"S": [("A1", 1), ("A2", 2), ("A3", "=SUM(A1:A2)")]})
        inject_cached_value(path, "xl/worksheets/sheet1.xml", "A3", "3")
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes())
        cell = [e for e in result["evidence"] if e.get("cell") == "A3"][0]
        self.assertEqual(cell["cached_value"], 3)
        self.assertEqual(result["counts"]["formulas_cached_unknown"], 0)

    def test_hidden_and_veryhidden_sheets_visited(self):
        path = _tmp("h.xlsx")
        make_xlsx(path,
                  sheets={"明": [("A1", "v")], "暗": [("A1", "h1")],
                          "深暗": [("A1", "h2")]},
                  hidden={"暗": "hidden", "深暗": "veryHidden"})
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["counts"]["sheets_total"], 3)
        self.assertEqual(result["counts"]["sheets_hidden"], 2)
        sheets = {e["sheet"]: e for e in result["evidence"]
                  if e["kind"] == "xlsx_sheet_info"}
        self.assertEqual(sheets["暗"]["sheet_state"], "hidden")
        self.assertEqual(sheets["深暗"]["sheet_state"], "veryHidden")
        locators = {e["locator"] for e in result["evidence"]
                    if e["kind"] == "xlsx_cell"}
        self.assertIn("暗!A1", locators)
        self.assertIn("深暗!A1", locators)

    def test_sparse_huge_sheet_scan_cap_partial(self):
        path = _tmp("sparse.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        ws["A1"] = "top"
        ws["A100000"] = "deep"  # 实际稀疏分布到十万行
        wb.save(path)
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes(),
                                   max_scan_rows=1000, max_scan_cells=1_000_000)
        self.assertEqual(result["status"], "partial")
        self.assertTrue(result["counts"]["scan_truncated"])
        self.assertEqual(result["counts"]["scanned_rows_total"], 1000)
        self.assertTrue(any("触顶" in n for n in result["notes"]))

    def test_wide_row_cell_cap_enforced_exactly(self):
        # 一行 20 格 > max_scan_cells=10：上限必须逐格生效，
        # scanned_cells_total 严格等于 10，status=partial（不得 success）
        path = _tmp("wide.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        from openpyxl.utils import get_column_letter

        for c in range(1, 21):
            ws.cell(row=1, column=c, value=f"v{c}")
        ws.cell(row=2, column=1, value="row2")  # 触顶后不应被扫描
        wb.save(path)
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes(), max_scan_cells=10)
        self.assertEqual(result["status"], "partial")
        self.assertTrue(result["counts"]["scan_truncated"])
        self.assertEqual(result["counts"]["scanned_cells_total"], 10)
        cells = [e for e in result["evidence"] if e["kind"] == "xlsx_cell"]
        self.assertFalse(any(e["cell"] == "A2" for e in cells))

    def test_row_cut_short_on_last_row_still_partial(self):
        # 只有 1 行 20 格、配额 10：行内截断发生在最后一行，while 正常退出，
        # 仍必须置 scan_truncated=True 并降 partial（右半行不得静默漏项）
        path = _tmp("lastrow.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "S"
        for c in range(1, 21):
            ws.cell(row=1, column=c, value=f"v{c}")
        wb.save(path)
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes(), max_scan_cells=10)
        self.assertEqual(result["status"], "partial")
        self.assertTrue(result["counts"]["scan_truncated"])
        self.assertEqual(result["counts"]["scanned_cells_total"], 10)
        self.assertEqual(result["counts"]["scanned_rows_total"], 1)
        self.assertTrue(any("触顶" in n for n in result["notes"]))

    def test_empty_workbook_fails(self):
        path = _tmp("e.xlsx")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb.create_sheet("空")
        wb.save(path)
        from biaojing import xlsx_parser

        result = xlsx_parser.parse(Path(path).read_bytes())
        self.assertEqual(result["status"], "failed")


class ZipContainerTests(unittest.TestCase):
    def test_bad_zip_rejected(self):
        data = b"this is not a zip at all" * 10
        with self.assertRaises(Exception):
            expand(data, max_entries=10, max_total_uncompressed=1 << 20)

    def test_entry_limit_whole_archive(self):
        path = _tmp("many.zip")
        with zipfile.ZipFile(path, "w") as zf:
            for i in range(5):
                zf.writestr(f"f{i}.txt", "x")
        data = Path(path).read_bytes()
        with self.assertRaises(Exception):
            expand(data, max_entries=3, max_total_uncompressed=1 << 20)

    def test_uncompressed_total_limit_whole_archive(self):
        path = _tmp("big.zip")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("a.bin", b"\x00" * 10_000)
        data = Path(path).read_bytes()
        with self.assertRaises(Exception):
            expand(data, max_entries=10, max_total_uncompressed=100)

    def test_path_traversal_entry_rejected(self):
        path = _tmp("trav.zip")
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("../evil.txt", "x")
            zf.writestr("ok.txt", "y")
        result = expand(Path(path).read_bytes(), 10, 1 << 20)
        self.assertEqual([e["name"] for e in result["entries"]], ["ok.txt"])
        self.assertIn("穿越", result["rejected"][0]["reason"])

    def test_duplicate_entry_names_both_kept(self):
        path = _tmp("dupname.zip")
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("same.txt", "one")
            zf.writestr("same.txt", "two")
        result = expand(buf.getvalue(), 10, 1 << 20)
        self.assertEqual(len(result["entries"]), 2)
        self.assertEqual(result["entries"][0]["data"], b"one")
        self.assertEqual(result["entries"][1]["data"], b"two")

    def test_nested_zip_entry_passthrough(self):
        path = _tmp("nest.zip")
        inner = io.BytesIO()
        with zipfile.ZipFile(inner, "w") as zf:
            zf.writestr("in.txt", "i")
        with zipfile.ZipFile(path, "w") as zf:
            zf.writestr("inner.zip", inner.getvalue())
        result = expand(Path(path).read_bytes(), 10, 1 << 20)
        self.assertEqual(len(result["entries"]), 1)
        self.assertEqual(result["entries"][0]["name"], "inner.zip")

    def test_guard_ooxml_limits(self):
        path = _tmp("g.xlsx")
        make_xlsx(path, sheets={"S": [("A1", 1)]})
        data = Path(path).read_bytes()
        self.assertIsNone(guard_ooxml(data, 100, 1 << 20))
        self.assertIsNotNone(guard_ooxml(data, 100, 10))       # 总量超限
        self.assertIsNotNone(guard_ooxml(data, 3, 1 << 20))    # 成员数超限


class CliEndToEndTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="biaojing_e2e_")

    def _p(self, name):
        return os.path.join(self.dir, name)

    def test_duplicates_kept_and_marked(self):
        a, b = self._p("a.pdf"), self._p("b_copy.pdf")
        make_pdf(a, ["同字节"])
        shutil.copy(a, b)
        out = self._p("out.json")
        payload = run_cli(a, b, output=out)
        recs = payload["files"]
        self.assertEqual(len(recs), 2)
        self.assertEqual(recs[0]["extract_status"], "success")
        self.assertEqual(recs[1]["extract_status"], "duplicate")
        self.assertEqual(recs[1]["duplicate_of"], recs[0]["file"])
        self.assertEqual(recs[0]["sha256"], recs[1]["sha256"])
        self.assertEqual(payload["summary"]["by_status"]["duplicate"], 1)

    def test_folder_expansion(self):
        sub = os.path.join(self.dir, "sub")
        os.makedirs(sub)
        make_pdf(os.path.join(sub, "1.pdf"), ["甲"])
        make_docx(os.path.join(sub, "2.docx"), ["乙"])
        out = self._p("out.json")
        payload = run_cli(self.dir, output=out)
        self.assertEqual(payload["summary"]["total_inputs"], 2)
        types = payload["summary"]["by_doc_type"]
        self.assertEqual(types.get("pdf"), 1)
        self.assertEqual(types.get("docx"), 1)

    def test_legacy_and_unknown_types(self):
        docp = self._p("old.doc")
        with open(docp, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0legacy")
        xls = self._p("old.xls")
        with open(xls, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0xls")
        unk = self._p("mystery.xyz")
        with open(unk, "wb") as f:
            f.write(b"\x00\x01unknown bits")
        out = self._p("out.json")
        payload = run_cli(docp, xls, unk, output=out)
        statuses = {r["doc_type"]: r["extract_status"] for r in payload["files"]}
        self.assertEqual(statuses["doc_legacy"], "pending_convert")
        self.assertEqual(statuses["xls_legacy"], "pending_convert")
        self.assertEqual(statuses["unknown"], "unknown")

    def test_bad_zip_and_isolated_failures_batch_survives(self):
        bad = self._p("bad.zip")
        with open(bad, "wb") as f:
            f.write(b"PK\x03\x04not really a zip")
        good = self._p("good.pdf")
        make_pdf(good, ["好文件"])
        out = self._p("out.json")
        payload = run_cli(bad, good, output=out)
        recs = {r["doc_type"]: r for r in payload["files"]}
        self.assertEqual(recs["zip"]["extract_status"], "failed")
        self.assertEqual(recs["pdf"]["extract_status"], "success")
        self.assertEqual(recs["zip"]["doc_type"], "zip")

    def test_zip_import_with_limits_and_traversal(self):
        archive = self._p("pack.zip")
        inner_docx = self._p("_inner.docx")
        make_docx(inner_docx, ["包内文档"])
        with zipfile.ZipFile(archive, "w") as zf:
            zf.write(inner_docx, "docs/包内文档.docx")
            zf.writestr("../逃逸.txt", "x")
        out = self._p("out.json")
        payload = run_cli(archive, output=out)
        recs = payload["files"]
        kinds = {}
        for r in recs:
            kinds[(r["origin"]["kind"], r["extract_status"])] = r
        entry_ok = [r for r in recs
                    if r["origin"]["kind"] == "zip_entry"
                    and r["extract_status"] == "success"]
        self.assertEqual(len(entry_ok), 1)
        self.assertTrue(entry_ok[0]["file"].endswith("::docs/包内文档.docx"))
        rejected = [r for r in recs
                    if r["origin"]["kind"] == "zip_entry_rejected"]
        self.assertEqual(len(rejected), 1)
        self.assertIn("穿越", rejected[0]["error"])

    def test_zip_entry_limit_cli(self):
        archive = self._p("over.zip")
        with zipfile.ZipFile(archive, "w") as zf:
            for i in range(5):
                zf.writestr(f"f{i}.txt", "x")
        out = self._p("out.json")
        payload = run_cli(archive, max_zip_entries=3, output=out)
        rec = payload["files"][0]
        self.assertEqual(rec["extract_status"], "failed")
        self.assertIn("超过限额", rec["error"])

    def test_ooxml_guard_via_cli(self):
        x = self._p("small.xlsx")
        make_xlsx(x, sheets={"S": [("A1", 1)]})
        out = self._p("out.json")
        payload = run_cli(x, max_zip_total_bytes=50, output=out)
        rec = payload["files"][0]
        self.assertEqual(rec["extract_status"], "failed")
        self.assertIn("解压总量", rec["error"])

    def test_sha256_in_output_matches_source_bytes(self):
        import hashlib
        p = self._p("h.pdf")
        make_pdf(p, ["哈希"])
        out = self._p("out.json")
        payload = run_cli(p, output=out)
        raw = Path(p).read_bytes()
        self.assertEqual(payload["files"][0]["sha256"],
                         hashlib.sha256(raw).hexdigest())

    def test_product_name_and_module_fields(self):
        p = self._p("n.pdf")
        make_pdf(p, ["字段"])
        out = self._p("out.json")
        payload = run_cli(p, output=out)
        self.assertEqual(payload["product_name"], "标镜")
        self.assertEqual(payload["product"], "标镜")
        self.assertEqual(payload["module"], "biaojing")

    def test_nonexistent_path_is_isolated_failure(self):
        out = self._p("out.json")
        payload = run_cli(os.path.join(self.dir, "不存在.pdf"), output=out)
        rec = payload["files"][0]
        self.assertEqual(rec["extract_status"], "failed")


class WalkerTests(unittest.TestCase):
    def test_zip_entry_reference_format(self):
        d = tempfile.mkdtemp(prefix="biaojing_w_")
        archive = os.path.join(d, "a.zip")
        with zipfile.ZipFile(archive, "w") as zf:
            zf.writestr("x.txt", "hello")
        refs = []
        for display, origin, data, error in iter_inputs([archive], 10, 1 << 20):
            refs.append((display, origin["kind"], data, error))
        self.assertEqual(len(refs), 1)
        display, kind, data, error = refs[0]
        self.assertTrue(display.startswith(archive + "::"))
        self.assertEqual(kind, "zip_entry")
        self.assertEqual(data, b"hello")
        self.assertIsNone(error)


if __name__ == "__main__":
    unittest.main()
