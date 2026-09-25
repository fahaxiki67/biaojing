# -*- coding: utf-8 -*-
"""长任务结果归一回归：后台任务直返形态不得被 UI 丢字段。

背景缺陷（UI 批量验收确认，诊断经独立复核）：
  - /api/upload 的长 PDF（≥3 页）与含图 DOCX 后台任务把 wb.ingest_bytes(...)
    的返回值原样作为 job.result——顶层 status/ref/candidates，无
    results/result 包裹键；
  - 页面 uploadFiles 只认 j.results / j.result 两个键 → 直接形态的
    status/ref/candidates 全部丢失，上传日志打出 "[ ? ]"，且业务上
    failed/rejected/partial 的文件在"已完成"任务里也被计成已处理。
本模块两层锁住行为（仅合成资料）：
  1) Python HTTP 全链路：钉住各 worker 的 job.result 契约（UI 必须兼容
     的形态），含取消路径；
  2) Node + DOM stub 实际执行页面 uploadFiles：直接形态必须打出真实
     状态（不得 "[ ? ]"），done/fail 统计与业务状态同源。
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from unittest.mock import patch

import docx as docx_lib
import pymupdf

from biaojing import pdf_parser, webapp, workspace


def wait_terminal(srv, job_id: str, timeout: float = 10) -> dict:
    base = f"http://127.0.0.1:{srv.server_address[1]}"
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        with urllib.request.urlopen(base + "/api/jobs/" + job_id) as response:
            job = json.load(response)
        if job["status"] not in ("queued", "running"):
            return job
        time.sleep(0.02)
    raise AssertionError("长任务超时未结束")


class WorkerJobResultContractTests(unittest.TestCase):
    """钉住各 worker 的 job.result 形态——UI 兼容层必须依据这些契约。"""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="biaojing_async_norm_")
        self.wb = workspace.Workbench(self.temp.name)
        self.srv = webapp.WorkbenchServer(("127.0.0.1", 0), self.wb)
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}"
        self.thread = threading.Thread(target=self.srv.serve_forever,
                                       daemon=True)
        self.thread.start()

    def tearDown(self):
        self.srv.shutdown()
        self.srv.server_close()
        self.wb.close()
        self.temp.cleanup()

    def upload(self, name: str, data: bytes) -> dict:
        req = urllib.request.Request(
            self.base + "/api/upload?name=" + urllib.parse.quote(name, safe=""),
            data=data,
            headers={"Origin": self.base,
                     "Content-Type": "application/octet-stream"})
        with urllib.request.urlopen(req, timeout=5) as response:
            return json.load(response)

    @staticmethod
    def pdf_bytes(pages: int = 3, text: bool = True) -> bytes:
        doc = pymupdf.open()
        for index in range(pages):
            page = doc.new_page()
            if text:
                page.insert_text((40, 60), f"合成文本页 {index}")
        data = doc.tobytes()
        doc.close()
        return data

    def test_long_pdf_job_result_is_single_item_with_true_status(self):
        # ≥3 页 PDF 走后台任务；job.result 必须是顶层单条结果（含真实
        # status/ref/candidates），页面据此打出真实状态
        started = self.upload("长文本件.pdf", self.pdf_bytes(3))
        self.assertIn("job_id", started)
        job = wait_terminal(self.srv, started["job_id"])
        self.assertEqual(job["status"], "completed")
        result = job["result"]
        self.assertNotIn("results", result)   # 直接形态：无列表包裹
        self.assertNotIn("result", result)    # 也无信封包裹
        self.assertEqual(result["status"], "success")
        self.assertIn("长文本件.pdf", result["ref"])
        self.assertEqual(len(result["sha256"]), 64)
        self.assertIsInstance(result["candidates"], int)
        self.assertIsInstance(result["counts"], dict)
        self.assertIsInstance(result["cancelled"], bool)
        self.assertEqual(result["counts"]["pages_total"], 3)

    def test_docx_media_job_result_is_single_item(self):
        image = pymupdf.Pixmap(
            pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), 0).tobytes("png")
        doc = docx_lib.Document()
        doc.add_paragraph().add_run().add_picture(io.BytesIO(image))
        stream = io.BytesIO()
        doc.save(stream)
        with patch.object(pdf_parser, "_ocr_runtime",
                          return_value=("tesseract", "chi_sim+eng", None)), \
                patch.object(pdf_parser, "ocr_image",
                             side_effect=lambda *a, **k: ("合成图片文字", None)):
            started = self.upload("含图Word.docx", stream.getvalue())
            self.assertIn("job_id", started)
            # 必须在 patch 上下文内等到任务结束：worker 线程在后台
            # 继续 OCR，提前退出 patch 会让真实 OCR 接管
            job = wait_terminal(self.srv, started["job_id"])
        self.assertEqual(job["status"], "completed")
        result = job["result"]
        self.assertNotIn("results", result)
        self.assertNotIn("result", result)
        self.assertEqual(result["status"], "success")
        self.assertIn("含图Word.docx", result["ref"])
        self.assertIsInstance(result["candidates"], int)
        self.assertIsInstance(result["cancelled"], bool)

    def test_zip_job_result_is_item_list(self):
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, "w") as bundle:
            bundle.writestr("member.pdf", self.pdf_bytes(1))
        started = self.upload("合成归档.zip", archive.getvalue())
        self.assertIn("job_id", started)
        job = wait_terminal(self.srv, started["job_id"])
        self.assertEqual(job["status"], "completed")
        result = job["result"]
        self.assertIsInstance(result["results"], list)
        self.assertIsInstance(result["cancelled"], bool)
        self.assertFalse(result["cancelled"])
        statuses = [item["status"] for item in result["results"]]
        self.assertIn("archived", statuses)    # 容器本体
        self.assertIn("success", statuses)     # 展开成员
        for item in result["results"]:
            self.assertIn("status", item)
            self.assertIn("ref", item)

    def test_cancelled_pdf_job_result_keeps_partial_status_and_pages(self):
        # 取消/中断口径不回归：已完成页保留、未处理页待处理；
        # job.result 顶层如实给出 cancelled/partial 与逐页计数
        started_event, release = threading.Event(), threading.Event()

        def slow_ocr(*args, **kwargs):
            started_event.set()
            self.assertTrue(release.wait(3), "test OCR was not released")
            return "扫描页面内容", None

        with patch.object(pymupdf.Page, "get_images", return_value=[(1,)]), \
                patch.object(pdf_parser, "_ocr_runtime",
                             return_value=("tesseract", "chi_sim", None)), \
                patch.object(pdf_parser, "_ocr_page", side_effect=slow_ocr):
            started = self.upload("扫描件.pdf", self.pdf_bytes(3, text=False))
            self.assertIn("job_id", started)
            self.assertTrue(started_event.wait(3), "background OCR did not start")
            req = urllib.request.Request(
                self.base + "/api/jobs/" + started["job_id"] + "/cancel",
                data=b"{}",
                headers={"Origin": self.base,
                         "Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=1) as response:
                self.assertTrue(json.load(response)["cancel_requested"])
            release.set()
            job = wait_terminal(self.srv, started["job_id"])
        self.assertEqual(job["status"], "cancelled")
        result = job["result"]
        self.assertTrue(result["cancelled"])
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["counts"]["pages_ocr"], 1)
        self.assertEqual(result["counts"]["pages_pending_ocr"], 2)
        pages = self.wb.conn.execute(
            "SELECT page_status FROM evidence_store ORDER BY page_no"
        ).fetchall()
        self.assertEqual([row["page_status"] for row in pages],
                         ["ocr", "pending_ocr", "pending_ocr"])


class UploadLogBehaviorTests(unittest.TestCase):
    """Node + DOM stub 实际执行页面 uploadFiles 的日志与计数回归。

    缺陷形态：任务轮询拿到 job.result 后只读 j.results / j.result 两个
    键——PDF/DOCX 直接返回形态（顶层 status/ref/candidates）全部丢失，
    日志打出 "[ ? ]" 并把业务失败计成已处理。修复后必须：
      a. 直接形态打出真实 status/ref/candidates（不得 "[ ? ]"）；
      b. failed/rejected/partial 一律计入失败/拒收，不计已处理；
      c. ZIP 列表形态与同步信封形态（既有口径）不回归；
      d. 取消路径提示"剩余页已保留待处理"且逐条状态如实。
    """

    maxDiff = None

    @classmethod
    def setUpClass(cls):
        if not shutil.which("node"):
            raise unittest.SkipTest("node 不可用")
        from biaojing.ui_page import INDEX_HTML
        match = re.search(r"<script>\n(.*?)</script>", INDEX_HTML, re.S)
        page_js = match.group(1)
        # /api/state 返回完整合法 state：页面末尾的 loadState() 正常完成
        state_stub = json.dumps({
            "product_name": "标镜",
            "findings": [],
            "coverage": {
                "total_input_occurrences": 0, "unique_files": 0,
                "by_status": {},
                "status_order": ["success", "partial", "failed",
                                 "pending_ocr", "pending_convert",
                                 "duplicate", "rejected", "archived",
                                 "unknown"]},
            "source_rows": [], "source_total": 0, "source_offset": 0,
            "source_limit": 100, "source_has_more": False,
            "candidates": [], "candidate_total": 0, "candidate_offset": 0,
            "candidate_limit": 200, "candidate_has_more": False,
            "confirmation_count": 0, "confirmation_history": [],
            "confirmation_history_total": 0,
            "file_binding_history": [], "file_binding_history_total": 0,
            "last_screen_run": {"status": "not_run"}})
        sha = "a" * 64
        scenarios = [
            {"id": "pdf_partial_direct", "filename": "160页扫描件.pdf",
             "upload": {"ok": True, "job_id": "job-pdf"},
             "job": {"job_id": "job-pdf", "kind": "pdf_upload",
                     "status": "completed", "page": 160, "total_pages": 160,
                     "stage": "已完成",
                     "result": {"sha256": sha, "status": "partial",
                                "ref": "160页扫描件.pdf（900000 字节）",
                                "doc_type": "pdf", "candidates": 0,
                                "counts": {"pages_total": 160,
                                           "pages_ocr": 151,
                                           "pages_pending_ocr": 9},
                                "cancelled": False}}},
            {"id": "pdf_failed_direct", "filename": "损坏件.pdf",
             "upload": {"ok": True, "job_id": "job-pdf"},
             "job": {"job_id": "job-pdf", "kind": "pdf_upload",
                     "status": "completed", "page": 3, "total_pages": 3,
                     "stage": "已完成",
                     "result": {"sha256": sha, "status": "failed",
                                "ref": "损坏件.pdf（300 字节）",
                                "doc_type": "pdf", "candidates": 0,
                                "counts": {}, "cancelled": False}}},
            {"id": "pdf_success_direct", "filename": "标书正本.pdf",
             "upload": {"ok": True, "job_id": "job-pdf"},
             "job": {"job_id": "job-pdf", "kind": "pdf_upload",
                     "status": "completed", "page": 5, "total_pages": 5,
                     "stage": "已完成",
                     "result": {"sha256": sha, "status": "success",
                                "ref": "标书正本.pdf（12000 字节）",
                                "doc_type": "pdf", "candidates": 3,
                                "counts": {"pages_total": 5}, "cancelled": False}}},
            {"id": "zip_mixed", "filename": "合成归档.zip",
             "upload": {"ok": True, "job_id": "job-zip"},
             "job": {"job_id": "job-zip", "kind": "zip_upload",
                     "status": "completed", "page": 3, "total_pages": 3,
                     "stage": "已完成",
                     "result": {"results": [
                         {"sha256": sha, "status": "archived",
                          "ref": "合成归档.zip"},
                         {"sha256": "b" * 64, "status": "success",
                          "ref": "合成归档.zip::a.pdf", "doc_type": "pdf",
                          "candidates": 1, "counts": {}, "cancelled": False},
                         {"ref": "合成归档.zip::huge.exe",
                          "status": "rejected", "reason": "条目类型被拒收"}],
                         "cancelled": False}}},
            {"id": "sync_success", "filename": "a.pdf",
             "upload": {"ok": True,
                        "result": {"sha256": sha, "status": "success",
                                   "ref": "a.pdf（10 字节）",
                                   "doc_type": "pdf", "candidates": 2,
                                   "counts": {}, "cancelled": False}},
             "job": None},
            {"id": "cancel_partial_direct", "filename": "扫描件.pdf",
             "upload": {"ok": True, "job_id": "job-pdf"},
             "job": {"job_id": "job-pdf", "kind": "pdf_upload",
                     "status": "cancelled", "page": 1, "total_pages": 3,
                     "stage": "已取消，未处理页面已保留",
                     "result": {"sha256": sha, "status": "partial",
                                "ref": "扫描件.pdf（800 字节）",
                                "doc_type": "pdf", "candidates": 0,
                                "counts": {"pages_total": 3, "pages_ocr": 1,
                                           "pages_pending_ocr": 2},
                                "cancelled": True}}},
        ]
        harness = """
const src=%s;
const STATE=%s;
const SCENARIOS=%s;
let CURRENT=null;
const els={};
function makeEl(){const n={_kids:[],_handlers:{},textContent:"",style:{},
  dataset:{},value:"",
  appendChild(c){n._kids.push(c);return c;},
  insertBefore(c){n._kids.unshift(c);return c;},
  addEventListener(t,f){n._handlers[t]=f;},remove(){},
  classList:{add(){},remove(){}},click(){}};
  return n;}
global.document={querySelector:s=>(els[s]=els[s]||makeEl()),
  createElement:()=>makeEl(),
  createTextNode:t=>({textContent:t,_kids:[]})};
global.window={open(){}};
global.setTimeout=()=>0;
global.fetch=async(url,opt)=>{
  if(url.indexOf("/api/state")===0)
    return {ok:true,status:200,json:async()=>JSON.parse(JSON.stringify(STATE))};
  if(url.indexOf("/api/upload?")===0)
    return {ok:true,status:200,json:async()=>CURRENT.upload};
  if(url.indexOf("/api/jobs/")===0)
    return {ok:true,status:200,json:async()=>CURRENT.job};
  throw new Error("unexpected fetch: "+url)};
// 真实分隔符用 String.fromCharCode(10) 生成，避免多层反斜杠转义歧义；
// eval 内同时覆盖 loadState，消除与页面自身调用的竞态
const NL=String.fromCharCode(10);
eval(src+NL+";globalThis.__T={uploadFiles,els};"
  + NL+"loadState=async()=>{};");
""" % (json.dumps(page_js), state_stub, json.dumps(scenarios))
        harness += """
(async()=>{
  const T=globalThis.__T;
  const out=[];
  for(const sc of SCENARIOS){
    CURRENT=sc;
    await T.uploadFiles([{f:{name:sc.filename,size:10},rel:sc.filename}]);
    out.push({id:sc.id, log:T.els["#uploadLog"].textContent});  }
  console.log(JSON.stringify(out));
})().catch(e=>{console.error("HARNESS_FAIL",e&&e.stack||e&&e.message);
  process.exit(1)});
"""
        cls.tmp = tempfile.mkdtemp(prefix="biaojing_async_norm_js_")
        cls.js_path = os.path.join(cls.tmp, "page.js")
        Path(cls.js_path).write_text(harness, encoding="utf-8")
        cls.proc = subprocess.run(
            ["node", cls.js_path], capture_output=True, text=True, timeout=60)

    def test_harness_ran(self):
        self.assertNotIn("HARNESS_FAIL", self.proc.stderr,
                         self.proc.stderr[:400])

    @classmethod
    def logs(cls) -> dict:
        if not hasattr(cls, "_logs"):
            cls._logs = {item["id"]: item["log"]
                         for item in json.loads(cls.proc.stdout)}
        return cls._logs

    # ---- a. 直接返回形态不丢字段：真实状态而非 "[ ? ]" ----

    def test_partial_direct_shape_logs_true_status_not_question_mark(self):
        log = self.logs()["pdf_partial_direct"]
        self.assertIn("[partial] 160页扫描件.pdf（900000 字节）", log)
        self.assertNotIn("[ ? ]", log)
        self.assertNotIn("[?]", log)
        # 待 OCR 页数通过加载状态可见即可；日志至少不丢状态与引用

    def test_failed_direct_shape_logs_true_status(self):
        log = self.logs()["pdf_failed_direct"]
        self.assertIn("[failed] 损坏件.pdf（300 字节）", log)
        self.assertNotIn("[ ? ]", log)

    def test_success_direct_shape_logs_status_ref_and_candidates(self):
        log = self.logs()["pdf_success_direct"]
        self.assertIn("[success] 标书正本.pdf（12000 字节）", log)
        self.assertIn("候选 3 条", log)
        self.assertNotIn("[ ? ]", log)

    # ---- b. done/fail 统计与业务状态同源 ----

    def test_partial_counts_as_not_done(self):
        log = self.logs()["pdf_partial_direct"]
        self.assertIn("上传结束：已处理 0，失败/拒收 1（含枚举失败 0），发现文件 1", log)

    def test_failed_counts_as_not_done(self):
        log = self.logs()["pdf_failed_direct"]
        self.assertIn("上传结束：已处理 0，失败/拒收 1（含枚举失败 0），发现文件 1", log)

    def test_success_counts_as_done(self):
        log = self.logs()["pdf_success_direct"]
        self.assertIn("上传结束：已处理 1，失败/拒收 0（含枚举失败 0），发现文件 1", log)

    # ---- c. ZIP 列表形态与同步信封形态（既有口径）不回归 ----

    def test_zip_item_list_shape_still_per_entry_and_counts(self):
        log = self.logs()["zip_mixed"]
        self.assertIn("[archived] 合成归档.zip", log)
        self.assertIn("[success] 合成归档.zip::a.pdf", log)
        self.assertIn("候选 1 条", log)
        self.assertIn("[rejected] 合成归档.zip::huge.exe", log)
        # rejected 条目使整批不计为已处理
        self.assertIn("上传结束：已处理 0，失败/拒收 1（含枚举失败 0），发现文件 1", log)

    def test_sync_envelope_shape_still_works(self):
        log = self.logs()["sync_success"]
        self.assertIn("[success] a.pdf（10 字节）", log)
        self.assertIn("候选 2 条", log)
        self.assertIn("上传结束：已处理 1，失败/拒收 0（含枚举失败 0），发现文件 1", log)

    # ---- d. 取消路径口径不回归且状态如实 ----

    def test_cancelled_job_keeps_preserved_pages_wording_and_true_status(self):
        log = self.logs()["cancel_partial_direct"]
        self.assertIn("[已取消] 扫描件.pdf：剩余页已保留待处理，可稍后重试", log)
        self.assertIn("[partial] 扫描件.pdf（800 字节）", log)
        self.assertNotIn("[ ? ]", log)
        # partial 不得计为已处理
        self.assertIn("上传结束：已处理 0，失败/拒收 1（含枚举失败 0），发现文件 1", log)


class ZeroByteUploadCoverageTests(unittest.TestCase):
    """D-01 回归：空文件上传被安全拒收（400）的同时必须计入覆盖率。

    缺陷：webapp._read_body 只在超限分支记录拒收，空体分支直接 400——
    该提交从覆盖率分母消失。修复口径：拒收语义不变（不得为空文件
    产出 artifact/哈希/证据），但覆盖率必须出现对应拒收记录。
    """

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(
            prefix="biaojing_zero_byte_")
        self.wb = workspace.Workbench(self.temp.name)
        self.srv = webapp.WorkbenchServer(("127.0.0.1", 0), self.wb)
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}"
        self.thread = threading.Thread(target=self.srv.serve_forever,
                                       daemon=True)
        self.thread.start()

    def tearDown(self):
        self.srv.shutdown()
        self.srv.server_close()
        self.wb.close()
        self.temp.cleanup()

    def test_empty_upload_is_rejected_and_counted_in_coverage(self):
        req = urllib.request.Request(
            self.base + "/api/upload?name="
            + urllib.parse.quote("空文件.pdf", safe=""),
            data=b"",
            headers={"Origin": self.base,
                     "Content-Type": "application/pdf"})
        try:
            urllib.request.urlopen(req, timeout=5)
            self.fail("空文件上传必须被 HTTP 400 拒收")
        except urllib.error.HTTPError as exc:
            with exc:  # 关闭错误响应句柄，避免 teardown ResourceWarning
                self.assertEqual(exc.code, 400)
        # 拒收语义不变：空文件不产出任何 artifact / 哈希 / 证据
        self.assertEqual(self.wb.conn.execute(
            "SELECT COUNT(*) FROM sources").fetchone()[0], 0)
        self.assertEqual(self.wb.conn.execute(
            "SELECT COUNT(*) FROM evidence_store").fetchone()[0], 0)
        self.assertEqual(self.wb.conn.execute(
            "SELECT COUNT(*) FROM candidates").fetchone()[0], 0)
        self.assertFalse(os.listdir(os.path.join(self.wb.root, "files")))
        # 但覆盖率必须包含该提交：拒收记录可见，分母不漏
        with urllib.request.urlopen(
                self.base + "/api/state?candidate_offset=0"
                "&candidate_limit=50&source_offset=0&source_limit=50",
                timeout=5) as response:
            state = json.load(response)
        rows = [row for row in state["source_rows"]
                if "空文件.pdf" in row["ref"]]
        self.assertEqual(len(rows), 1, "空文件上传未进入覆盖率")
        self.assertEqual(rows[0]["status"], "rejected")
        coverage = state["coverage"]
        self.assertEqual(coverage["by_status"].get("rejected"), 1)
        self.assertEqual(coverage["total_input_occurrences"], 1)


if __name__ == "__main__":
    unittest.main()
