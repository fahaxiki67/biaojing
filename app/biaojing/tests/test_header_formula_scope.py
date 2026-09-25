# -*- coding: utf-8 -*-
"""标镜表头建议的公式行作用域回归（F 线声明的建议路径公式行限制）。

口径：
  - 公式单元格必须计入行结构（非空行集合/双层表头护栏/截断判定）——
    仅含公式单元格的行不得被当成整行全空的分隔行而静默截断表头作用域；
  - 公式串本身仍绝不作为建议值输出（缓存值语义由 price_lines 路径处理）。

运行：
    cd app && python3 -W error::ResourceWarning -m unittest \\
        biaojing.tests.test_header_formula_scope -v
"""

from __future__ import annotations

import io
import unittest

import openpyxl

from biaojing import candidates, xlsx_parser


def _evidence_and_suggestions():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "报价"
    ws.append(["投标人名称", "综合单价(元)", "清单名称", "规格型号"])
    ws.append(["甲公司", 100, "电线", "BV"])
    ws["B3"] = "=B2"      # 全公式行：无缓存（openpyxl 保存不含缓存值）
    ws["C3"] = "=C2"
    ws["D3"] = "=D2"
    ws.append(["乙公司", 999, "开关", "10A"])  # r4
    buf = io.BytesIO()
    wb.save(buf)
    rec = xlsx_parser.parse(buf.getvalue())
    tagged = candidates.assign_evidence_ids("sha", rec["evidence"])
    sugs = candidates._xlsx_header_suggestions(tagged)
    return tagged, sugs


class HeaderFormulaRowScopeTests(unittest.TestCase):

    def test_formula_row_does_not_truncate_scope(self):
        _, sugs = _evidence_and_suggestions()
        by_loc = {s["locator_display"]: s for s in sugs}
        # r4 在全公式行（r3）之后：作用域不得被公式行截断
        self.assertIn("报价!A4", by_loc, "乙公司（r4）应仍有 bidder_name 建议")
        self.assertEqual(by_loc["报价!A4"]["value"], "乙公司")
        self.assertIn("报价!B4", by_loc, "r4 单价建议不得因公式行丢失")
        self.assertEqual(by_loc["报价!B4"]["value"], 999)
        self.assertIn("报价!D4", by_loc)

    def test_formula_cells_never_become_suggestions(self):
        _, sugs = _evidence_and_suggestions()
        bad = [s for s in sugs
               if isinstance(s.get("value"), str)
               and s["value"].startswith("=")]
        self.assertEqual(bad, [], f"公式串漏为建议值：{bad[:3]}")
        locs = {s["locator_display"] for s in sugs}
        for cell in ("报价!B3", "报价!C3", "报价!D3"):
            self.assertNotIn(cell, locs, f"{cell} 是公式串，不得出建议")


if __name__ == "__main__":
    unittest.main()
