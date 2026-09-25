"""跨平台入口及失败隔离回归；仅使用临时合成资料。"""
import concurrent.futures
import io
import json
import os
import sqlite3
from pathlib import Path
import subprocess
import sys
import tempfile
import threading
import time
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch
import urllib.error
import urllib.request

import docx as docx_lib
import pymupdf
from biaojing import pdf_parser, workspace, webapp


def pdf_bytes():
    with pymupdf.open() as doc:
        doc.new_page().insert_text((50, 50), 'TEST 123')
        return doc.tobytes()


class CompatibilityTests(unittest.TestCase):
    def test_pdf_cancel_keeps_unprocessed_pages_retryable(self):
        with pymupdf.open() as doc:
            for _ in range(3):
                doc.new_page()
            data = doc.tobytes()
        stop = threading.Event()
        progress = []

        def finish_first_page(page, total, stage, status):
            progress.append((page, total, stage, status))
            if page == 1 and stage == 'page_done':
                stop.set()

        with patch.object(pymupdf.Page, 'get_images', return_value=[(1,)]), \
                patch.object(pdf_parser, '_ocr_runtime',
                             return_value=('tesseract', 'chi_sim', None)), \
                patch.object(pdf_parser, '_ocr_page',
                             return_value=('已识别页面文字', None)):
            result = pdf_parser.parse(data, progress=finish_first_page,
                                     cancelled=stop.is_set)

        self.assertTrue(result['cancelled'])
        self.assertEqual(result['counts'], {
            'pages_total': 3, 'pages_text': 0, 'pages_ocr': 1,
            'pages_pending_ocr': 2, 'pages_blank': 0})
        self.assertEqual([e['page_status'] for e in result['evidence']],
                         ['ocr', 'pending_ocr', 'pending_ocr'])
        self.assertTrue(all('尚未扫描' in e['page_error']
                            for e in result['evidence'][1:]))
        self.assertEqual([(p, st) for p, _, st, _ in progress],
                         [(1, 'processing'), (1, 'page_done')])

    def test_ocr_retry_only_updates_pending_pages_and_preserves_existing_facts(self):
        with pymupdf.open() as doc:
            doc.new_page().insert_text((50, 50), '项目名称：测试项目')
            doc.new_page()
            data = doc.tobytes()
        original_images = pymupdf.Page.get_images

        def images_for_scanned_page(page, full=False):
            return [(1,)] if page.number == 1 else original_images(page, full=full)

        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            try:
                with patch.object(pymupdf.Page, 'get_images', images_for_scanned_page), \
                        patch.object(pdf_parser, '_ocr_runtime',
                                     return_value=('tesseract', 'chi_sim', None)), \
                        patch.object(pdf_parser, '_ocr_page',
                                     return_value=('', '首次识别失败')):
                    imported = wb.ingest_bytes('two-pages.pdf', data)
                sha = imported['sha256']
                rows = wb.conn.execute(
                    'SELECT evidence_id, page_no, page_status FROM evidence_store'
                    ' WHERE sha256=? ORDER BY page_no', (sha,)).fetchall()
                self.assertEqual([r['page_status'] for r in rows], ['text', 'pending_ocr'])
                ids_before = [r['evidence_id'] for r in rows]
                wb.conn.execute(
                    "UPDATE evidence_store SET extract_version='legacy-test'"
                    ' WHERE sha256=? AND page_no=1', (sha,))
                wb.conn.commit()
                confirmed = wb.confirm_field('EV-RETRY', 'LOT-RETRY', 'BID-RETRY',
                                             'project_name', '测试项目', ids_before[0])
                self.assertTrue(confirmed['ok'], confirmed)

                calls = []
                def retry_page(page, *args):
                    calls.append(page.number + 1)
                    return '投标人名称：恢复测试公司', '已通过稀疏文字模式补试，需对照原页核实'

                with patch.object(pymupdf.Page, 'get_images', images_for_scanned_page), \
                        patch.object(pdf_parser, '_ocr_runtime',
                                     return_value=('tesseract', 'chi_sim', None)), \
                        patch.object(pdf_parser, '_ocr_page', side_effect=retry_page):
                    result = wb.retry_ocr(sha)

                rows_after = wb.conn.execute(
                    'SELECT evidence_id, page_no, page_status, extract_version'
                    ' FROM evidence_store WHERE sha256=? ORDER BY page_no', (sha,)).fetchall()
                self.assertEqual(calls, [2])
                self.assertEqual([r['evidence_id'] for r in rows_after], ids_before)
                self.assertEqual([r['page_status'] for r in rows_after], ['text', 'ocr'])
                self.assertEqual(rows_after[0]['extract_version'], 'legacy-test')
                self.assertEqual(rows_after[1]['extract_version'], pdf_parser.EXTRACTION_VERSION)
                recovered_page = wb.evidence_by_id(ids_before[1])
                self.assertIn('稀疏文字模式', recovered_page['page_note'])
                self.assertEqual(result['counts'], {
                    'pages_total': 2, 'pages_text': 1, 'pages_ocr': 1,
                    'pages_pending_ocr': 0, 'pages_blank': 0})
                self.assertEqual(wb.state()['confirmation_count'], 1)
                self.assertEqual(wb.state()['source_rows'][0]['status'], 'success')
                self.assertEqual(wb.original_bytes(sha), data)
                recovered = [c for c in wb.state()['candidates']
                             if c['field'] == 'bidder_name']
                self.assertEqual(len(recovered), 1)
                self.assertEqual(recovered[0]['evidence_id'], ids_before[1])
                self.assertIn('第 2 页 已通过稀疏文字模式',
                              wb.state()['source_rows'][0]['reason'])
            finally:
                wb.close()

    def test_legacy_pdf_evidence_migrates_as_pending_with_page_number(self):
        with tempfile.TemporaryDirectory() as root:
            db = sqlite3.connect(Path(root) / 'biaojing.sqlite3')
            db.executescript('''
                CREATE TABLE sources(sha256 TEXT PRIMARY KEY, doc_type TEXT,
                    status TEXT, parser TEXT, counts_json TEXT, metadata_json TEXT,
                    source_type TEXT, first_ref TEXT, imported_at TEXT);
                CREATE TABLE evidence_store(evidence_id TEXT PRIMARY KEY, sha256 TEXT,
                    ordinal INTEGER, kind TEXT, locator_display TEXT, locator_json TEXT,
                    quote TEXT, source_type TEXT, ocr_used INTEGER NOT NULL DEFAULT 0);
            ''')
            sha = 'a' * 64
            db.execute('INSERT INTO sources VALUES(?,?,?,?,?,?,?,?,?)',
                       (sha, 'pdf', 'pending_ocr', 'pymupdf', '{}', '{}',
                        'unknown', 'old.pdf', '2026-09-25'))
            db.execute('INSERT INTO evidence_store VALUES(?,?,?,?,?,?,?,?,?)',
                       ('E-legacy-0002', sha, 2, 'pdf_page', 'page 2',
                        '{"kind":"pdf_page","page":2}', '', 'unknown', 0))
            db.commit()
            db.close()

            wb = workspace.Workbench(root)
            try:
                row = wb.conn.execute(
                    'SELECT page_no, page_status, extract_version FROM evidence_store'
                ).fetchone()
                self.assertEqual(row['page_no'], 2)
                self.assertEqual(row['page_status'], 'pending_ocr')
                self.assertIn('legacy', row['extract_version'])
                self.assertIn('extract_version',
                              {r[1] for r in wb.conn.execute('PRAGMA table_info(sources)')})
            finally:
                wb.close()

    def test_empty_ocr_retries_sparse_but_short_noise_stays_pending(self):
        page = SimpleNamespace(rect=SimpleNamespace(width=100, height=100),
                               get_pixmap=Mock(return_value=SimpleNamespace(tobytes=lambda _: b'png')))
        for text, accepted in [('利润表' * 10, True), ('abc', False)]:
            responses = [subprocess.CompletedProcess([], 0, b'', b''),
                         subprocess.CompletedProcess([], 0, text.encode(), b'')]
            with patch.object(pdf_parser.subprocess, 'run', side_effect=responses) as run:
                value, note = pdf_parser._ocr_page(page, 'tesseract', 'chi_sim', 15)
            self.assertEqual(bool(value), accepted)
            self.assertIn('稀疏', note)
            self.assertEqual([c.args[0][-1] for c in run.call_args_list], ['3', '11'])
            self.assertLessEqual(run.call_args_list[1].kwargs['timeout'],
                                 run.call_args_list[0].kwargs['timeout'])

    def test_workspace_guard_rejects_double_open_and_releases(self):
        with tempfile.TemporaryDirectory() as root:
            first = webapp.workspace_guard(root)
            try:
                with self.assertRaises(sqlite3.OperationalError):
                    webapp.workspace_guard(root)
            finally:
                first.close()
            second = webapp.workspace_guard(root)
            second.close()

    def test_launch_check_works_outside_project_with_minimal_path(self):
        launch = Path(__file__).resolve().parents[3] / 'launch.py'
        with tempfile.TemporaryDirectory(prefix='标镜 启动 ') as cwd:
            result = subprocess.run([sys.executable, str(launch), '--check'],
                                    cwd=cwd, env={**os.environ, 'PATH': '/usr/bin:/bin'},
                                    capture_output=True, text=True, encoding='utf-8', timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('文档解析依赖可加载', result.stdout)

    def test_missing_python_dependency_is_readable(self):
        launch = Path(__file__).resolve().parents[3] / 'launch.py'
        result = subprocess.run([sys.executable, '-S', str(launch), '--check'],
                                capture_output=True, text=True, encoding='utf-8', timeout=10)
        self.assertEqual(result.returncode, 2)
        self.assertIn('依赖 pymupdf 无法加载', result.stderr)

    def test_macos_finder_ocr_without_shell_path(self):
        with patch.object(sys, 'platform', 'darwin'), \
                patch.dict(os.environ, {}, clear=True), \
                patch.object(pdf_parser.shutil, 'which', return_value=None), \
                patch.object(os.path, 'isfile', side_effect=lambda p: p == '/opt/homebrew/bin/tesseract'), \
                patch.object(pdf_parser.subprocess, 'run', return_value=
                             subprocess.CompletedProcess([], 0, 'chi_sim\neng\n', '')):
            self.assertEqual(pdf_parser._ocr_runtime()[0], '/opt/homebrew/bin/tesseract')

    def test_coverage_headers_match_status_order(self):
        import re
        from biaojing.ui_page import INDEX_HTML
        header = INDEX_HTML.split('id="coverage"', 1)[1].split('</thead>', 1)[0]
        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            try:
                self.assertEqual(len(re.findall(r'<th>', header)),
                                 len(wb.coverage()['status_order']) + 1)
            finally:
                wb.close()

    def test_confirmation_wrong_id_type_rejected_without_write(self):
        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            try:
                result = wb.confirm_field(['bad'], 'lot', 'bid', 'total_price',
                                          10, None, action='unknown')
                self.assertFalse(result['ok'])
                self.assertEqual(wb.state()['confirmation_count'], 0)
            finally:
                wb.close()

    def test_windows_ocr_discovery_without_path(self):
        with tempfile.TemporaryDirectory() as root:
            binary = Path(root) / 'Tesseract-OCR' / 'tesseract.exe'
            binary.parent.mkdir()
            binary.touch()
            with patch.object(sys, 'platform', 'win32'), \
                    patch.dict(os.environ, {'ProgramFiles': root}, clear=True), \
                    patch.object(pdf_parser.shutil, 'which', return_value=None), \
                    patch.object(pdf_parser.subprocess, 'run', return_value=
                                 subprocess.CompletedProcess([], 0, 'chi_sim\neng\n', '')):
                found, lang, error = pdf_parser._ocr_runtime()
            self.assertEqual(found, str(binary))
            self.assertEqual(lang, 'chi_sim+eng')
            self.assertIsNone(error)

    def test_render_failure_is_page_failure(self):
        page = SimpleNamespace(rect=SimpleNamespace(width=100, height=100),
                               get_pixmap=Mock(side_effect=RuntimeError('bad image')))
        text, error = pdf_parser._ocr_page(page, 'tesseract', 'chi_sim', 15)
        self.assertEqual(text, '')
        self.assertIn('渲染', error)

    def test_bad_page_does_not_discard_later_pages(self):
        with pymupdf.open() as doc:
            doc.new_page().insert_text((50, 50), 'FIRST')
            doc.new_page().insert_text((50, 50), 'SECOND')
            data = doc.tobytes()
        original = pymupdf.Page.get_text
        def get_text(page, *args, **kwargs):
            if page.number == 0:
                raise RuntimeError('bad page')
            return original(page, *args, **kwargs)
        with patch.object(pymupdf.Page, 'get_text', get_text), \
                patch.object(pdf_parser, '_ocr_page', return_value=('', 'OCR 失败')):
            result = pdf_parser.parse(data)
        self.assertEqual(result['status'], 'partial')
        self.assertEqual(len(result['evidence']), 2)
        self.assertIn('SECOND', result['evidence'][1]['text'])
        self.assertNotEqual(result['evidence'][0]['page_status'], 'blank')

    def test_parallel_same_file_is_one_source_and_two_refs(self):
        data = pdf_bytes()
        original = pdf_parser.parse
        barrier = threading.Barrier(2)
        def slow_parse(raw):
            time.sleep(0.05)
            return original(raw)
        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            try:
                def ingest(name):
                    barrier.wait(timeout=5)
                    return wb.ingest_bytes(name, data)['status']
                with patch.object(pdf_parser, 'parse', slow_parse), \
                        concurrent.futures.ThreadPoolExecutor(2) as pool:
                    statuses = list(pool.map(ingest, ['a.pdf', 'b.pdf']))
                self.assertEqual(sorted(statuses), ['duplicate', 'success'])
                self.assertEqual(wb.state()['coverage']['unique_files'], 1)
                self.assertEqual(len(wb.state()['source_rows']), 2)
            finally:
                wb.close()

    def test_parse_failure_reason_survives_import(self):
        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            try:
                wb.ingest_bytes('broken.pdf', b'bad pdf')
                self.assertIn('PDF', wb.state()['source_rows'][0]['reason'] or '')
            finally:
                wb.close()

    def test_invalid_json_returns_readable_http_error(self):
        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            srv = webapp.WorkbenchServer(('127.0.0.1', 0), wb)
            thread = threading.Thread(target=srv.serve_forever)
            thread.start()
            base = 'http://127.0.0.1:' + str(srv.server_port)
            try:
                for data in [b'{', b'[]', b'null']:
                    req = urllib.request.Request(base + '/api/confirm', data=data,
                                                 headers={'Origin': base})
                    with self.assertRaises(urllib.error.HTTPError) as ctx:
                        urllib.request.urlopen(req, timeout=5)
                    with ctx.exception as response:
                        self.assertEqual(response.code, 400)
                        self.assertIn('error', json.load(response))
            finally:
                srv.shutdown()
                thread.join(timeout=5)
                srv.server_close()
                wb.close()

    def test_retry_ocr_route_validates_sha_and_calls_workbench(self):
        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            srv = webapp.WorkbenchServer(('127.0.0.1', 0), wb)
            thread = threading.Thread(target=srv.serve_forever)
            thread.start()
            base = 'http://127.0.0.1:' + str(srv.server_port)
            try:
                def post(sha):
                    req = urllib.request.Request(
                        base + '/api/retry_ocr',
                        data=json.dumps({'sha256': sha}).encode(),
                        headers={'Origin': base, 'Content-Type': 'application/json'})
                    try:
                        with urllib.request.urlopen(req, timeout=5) as response:
                            return response.status, json.load(response)
                    except urllib.error.HTTPError as exc:
                        with exc:
                            return exc.code, json.load(exc)

                code, result = post('bad')
                self.assertEqual(code, 400)
                self.assertIn('SHA-256', result['error'])
                with patch.object(wb, 'retry_ocr', return_value={'ok': True,
                                                                 'retried_pages': 1}) as retry:
                    code, result = post('a' * 64)
                self.assertEqual(code, 200)
                self.assertIn('job_id', result)
                job_id = result['job_id']
                deadline = time.monotonic() + 3
                while time.monotonic() < deadline:
                    with urllib.request.urlopen(base + '/api/jobs/' + job_id) as response:
                        job = json.load(response)
                    if job['status'] not in ('queued', 'running'):
                        break
                    time.sleep(.01)
                self.assertEqual(job['status'], 'completed')
                self.assertEqual(job['result']['retried_pages'], 1)
                retry.assert_called_once()
                self.assertEqual(retry.call_args.args, ('a' * 64,))
                self.assertIn('progress', retry.call_args.kwargs)
                self.assertIn('cancelled', retry.call_args.kwargs)
            finally:
                srv.shutdown()
                thread.join(timeout=5)
                srv.server_close()
                wb.close()

    def test_long_pdf_upload_progress_and_cancel_endpoint_stays_responsive(self):
        with pymupdf.open() as doc:
            for _ in range(3):
                doc.new_page()
            data = doc.tobytes()
        started, release = threading.Event(), threading.Event()

        def slow_ocr(*args):
            started.set()
            self.assertTrue(release.wait(3), 'test OCR was not released')
            return '扫描页面内容', None

        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            srv = webapp.WorkbenchServer(('127.0.0.1', 0), wb)
            thread = threading.Thread(target=srv.serve_forever, daemon=True)
            thread.start()
            base = 'http://127.0.0.1:' + str(srv.server_port)
            try:
                with patch.object(pymupdf.Page, 'get_images', return_value=[(1,)]), \
                        patch.object(pdf_parser, '_ocr_runtime',
                                     return_value=('tesseract', 'chi_sim', None)), \
                        patch.object(pdf_parser, '_ocr_page', side_effect=slow_ocr):
                    req = urllib.request.Request(
                        base + '/api/upload?name=long.pdf', data=data,
                        headers={'Origin': base, 'Content-Type': 'application/pdf'})
                    with urllib.request.urlopen(req, timeout=3) as response:
                        started_job = json.load(response)
                    self.assertIn('job_id', started_job)
                    job_id = started_job['job_id']
                    self.assertTrue(started.wait(3), 'background OCR did not start')
                    req = urllib.request.Request(
                        base + '/api/jobs/' + job_id + '/cancel', data=b'{}',
                        headers={'Origin': base, 'Content-Type': 'application/json'})
                    with urllib.request.urlopen(req, timeout=1) as response:
                        cancel_reply = json.load(response)
                    self.assertTrue(cancel_reply['cancel_requested'])
                    release.set()
                    deadline = time.monotonic() + 5
                    while time.monotonic() < deadline:
                        with urllib.request.urlopen(base + '/api/jobs/' + job_id) as response:
                            job = json.load(response)
                        if job['status'] not in ('queued', 'running'):
                            break
                        time.sleep(.01)

                self.assertEqual(job['status'], 'cancelled')
                self.assertEqual(job['result']['counts']['pages_ocr'], 1)
                self.assertEqual(job['result']['counts']['pages_pending_ocr'], 2)
                state = wb.state()
                self.assertEqual(state['source_rows'][0]['status'], 'partial')
                self.assertIn('用户取消处理', state['source_rows'][0]['reason'])
                pages = wb.conn.execute(
                    'SELECT page_no, page_status, page_error FROM evidence_store'
                    ' ORDER BY page_no').fetchall()
                self.assertEqual([row['page_status'] for row in pages],
                                 ['ocr', 'pending_ocr', 'pending_ocr'])
                self.assertTrue(all('尚未扫描' in row['page_error']
                                    for row in pages[1:]))
            finally:
                release.set()
                srv.shutdown()
                thread.join(timeout=5)
                srv.server_close()
                wb.close()

    def test_docx_image_upload_runs_background_and_cancels_at_image_boundary(self):
        image = pymupdf.Pixmap(
            pymupdf.csRGB, pymupdf.IRect(0, 0, 2, 2), 0).tobytes('png')
        doc = docx_lib.Document()
        for _ in range(2):
            doc.add_paragraph().add_run().add_picture(io.BytesIO(image))
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'images.docx'
            doc.save(path)
            data = path.read_bytes()
        started, release = threading.Event(), threading.Event()

        def slow_ocr(*args):
            started.set()
            self.assertTrue(release.wait(3), 'test OCR was not released')
            return '投标人名称：华夏建设有限公司', None

        with tempfile.TemporaryDirectory() as root:
            wb = workspace.Workbench(root)
            srv = webapp.WorkbenchServer(('127.0.0.1', 0), wb)
            thread = threading.Thread(target=srv.serve_forever, daemon=True)
            thread.start()
            base = 'http://127.0.0.1:' + str(srv.server_port)
            try:
                with patch.object(pdf_parser, '_ocr_runtime',
                                  return_value=('tesseract', 'chi_sim+eng', None)), \
                        patch.object(pdf_parser, 'ocr_image', side_effect=slow_ocr):
                    req = urllib.request.Request(
                        base + '/api/upload?name=images.docx', data=data,
                        headers={'Origin': base,
                                 'Content-Type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'})
                    with urllib.request.urlopen(req, timeout=3) as response:
                        started_job = json.load(response)
                    self.assertIn('job_id', started_job)
                    job_id = started_job['job_id']
                    self.assertTrue(started.wait(3), 'Word OCR did not start')
                    req = urllib.request.Request(
                        base + '/api/jobs/' + job_id + '/cancel', data=b'{}',
                        headers={'Origin': base, 'Content-Type': 'application/json'})
                    with urllib.request.urlopen(req, timeout=1) as response:
                        self.assertTrue(json.load(response)['cancel_requested'])
                    release.set()
                    deadline = time.monotonic() + 5
                    while time.monotonic() < deadline:
                        with urllib.request.urlopen(base + '/api/jobs/' + job_id) as response:
                            job = json.load(response)
                        if job['status'] not in ('queued', 'running'):
                            break
                        time.sleep(.01)
                self.assertEqual(job['status'], 'cancelled')
                self.assertTrue(job['result']['cancelled'])
                self.assertEqual(job['result']['counts']['images_pending_ocr'], 1)
                rows = wb.conn.execute(
                    "SELECT page_status, page_error FROM evidence_store"
                    " WHERE kind='docx_image' ORDER BY ordinal").fetchall()
                self.assertEqual([r['page_status'] for r in rows],
                                 ['ocr', 'pending_ocr'])
                self.assertIn('用户取消处理', rows[1]['page_error'])
            finally:
                release.set()
                srv.shutdown()
                thread.join(timeout=5)
                srv.server_close()
                wb.close()
