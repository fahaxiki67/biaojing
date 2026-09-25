# -*- coding: utf-8 -*-
"""标镜 XLSX 合并单元格披露回归（C 线审查 P2-8 解析层半边）。

口径：
  - 工作簿存在合并区域时，解析层逐区域输出 kind=xlsx_merged_range 的
    披露证据（locator 为 `表!区域`，quote 说明"仅左上角持有值，其余为空，
    候选层不自动填充，须人工确认"）；
  - 该披露不计入 cells_total，不改变 success/partial/failed 状态口径；
  - counts.merged_ranges_total 如实统计，并设硬上限防畸形文件；
  - 无合并区域时不产出该类证据（既有行为零变化）。
  合并区域到候选行/列的自动填充属于 candidates 层（另册处理），本层只
  负责把结构如实暴露给人工与下游。

运行：
    cd app && python3 -W error::ResourceWarning -m unittest \\
        biaojing.tests.test_xlsx_merged_disclosure -v
"""

from __future__ import annotations

import io
import unittest

import openpyxl

from biaojing import xlsx_parser


def _wb_with_merges(merges: list[str], anchor_values: dict | None = None,
                    title: str = "报价") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(["投标人", "综合单价(元)", "清单名称"])
    ws.append(["甲公司", 100, "电线"])
    for ref in merges:
        ws.merge_cells(ref)
    for coord, value in (anchor_values or {}).items():
        ws[coord] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class XlsxMergedDisclosureTests(unittest.TestCase):

    def test_vertical_merge_is_disclosed(self):
        data = _wb_with_merges(["A3:A5"], {"A3": "甲公司"})
        rec = xlsx_parser.parse(data)
        merged = [e for e in rec["evidence"]
                  if e["kind"] == "xlsx_merged_range"]
        self.assertEqual(len(merged), 1)
        item = merged[0]
        self.assertEqual(item["sheet"], "报价")
        self.assertEqual(item["range"], "A3:A5")
        self.assertEqual(item["locator"], "报价!A3:A5")
        self.assertIn("左上角", item["quote"])
        self.assertIn("人工", item["quote"])
        self.assertEqual(rec["counts"]["merged_ranges_total"], 1)
        self.assertEqual(rec["counts"]["merged_ranges_incomplete"], False)
        self.assertEqual(rec["status"], "success")

    def test_no_merge_no_disclosure(self):
        data = _wb_with_merges([])
        rec = xlsx_parser.parse(data)
        self.assertFalse(any(e["kind"] == "xlsx_merged_range"
                             for e in rec["evidence"]))
        self.assertEqual(rec["counts"]["merged_ranges_total"], 0)
        self.assertEqual(rec["status"], "success")

    def test_multiple_ranges_all_disclosed(self):
        data = _wb_with_merges(["A3:A4", "B3:B4"], {"A3": "甲公司"})
        rec = xlsx_parser.parse(data)
        ranges = sorted(e["range"] for e in rec["evidence"]
                        if e["kind"] == "xlsx_merged_range")
        self.assertEqual(ranges, ["A3:A4", "B3:B4"])
        self.assertEqual(rec["counts"]["merged_ranges_total"], 2)

    def test_disclosure_does_not_count_as_cells(self):
        with_merge = xlsx_parser.parse(_wb_with_merges(["A3:A4"]))
        without = xlsx_parser.parse(_wb_with_merges([]))
        self.assertEqual(with_merge["counts"]["cells_total"],
                         without["counts"]["cells_total"])

    def test_merge_cap_is_disclosed_not_silent(self):
        # 畸形超量合并区域：触顶必须留痕（note），并如实降为 partial
        merges = [f"C{r}:D{r}" for r in range(3, 3 + 1001)]
        data = _wb_with_merges(merges)
        rec = xlsx_parser.parse(data)
        merged = [e for e in rec["evidence"]
                  if e["kind"] == "xlsx_merged_range"]
        self.assertEqual(len(merged), 1000)
        self.assertTrue(any("上限" in n for n in rec["notes"]),
                        f"触顶未留痕：{rec['notes']}")
        self.assertEqual(rec["counts"]["merged_ranges_incomplete"], True)
        self.assertEqual(rec["status"], "partial")

    def test_multi_sheet_aggregate_cap(self):
        # 全局每工作簿上限：两表各 600 条（合计 1200 > 1000），
        # 第二表只收到 400 条即止，触顶如实降为 partial
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "一表"
        ws2 = wb.create_sheet("二表")
        for r in range(3, 603):
            ws1.merge_cells(f"A{r}:B{r}")
            ws2.merge_cells(f"A{r}:B{r}")
        buf = io.BytesIO()
        wb.save(buf)
        rec = xlsx_parser.parse(buf.getvalue())
        merged = [e for e in rec["evidence"]
                  if e["kind"] == "xlsx_merged_range"]
        self.assertEqual(len(merged), 1000)
        per_sheet = {e["sheet"] for e in merged}
        self.assertEqual(per_sheet, {"一表", "二表"})
        self.assertTrue(any("上限" in n for n in rec["notes"]))
        self.assertEqual(rec["counts"]["merged_ranges_incomplete"], True)
        self.assertEqual(rec["status"], "partial")

    def test_unreadable_merge_part_is_reported_not_silent(self):
        # 工作表部件损坏：辅助层不得崩溃，必须置 incomplete 并留痕
        import zipfile as _zipfile
        good = _wb_with_merges(["A2:A3"])
        src = {}
        with _zipfile.ZipFile(io.BytesIO(good)) as z:
            for name in z.namelist():
                src[name] = z.read(name)
        src["xl/worksheets/sheet1.xml"] = b"<worksheet><mergeCells><mergeCell"
        buf = io.BytesIO()
        with _zipfile.ZipFile(buf, "w") as z:
            for name, payload in src.items():
                z.writestr(name, payload)
        merges, notes, incomplete = xlsx_parser._merged_ranges_by_sheet(
            buf.getvalue())
        self.assertFalse(merges)
        self.assertTrue(incomplete)
        self.assertTrue(any("不完整" in n for n in notes),
                        f"损坏部件未留痕：{notes}")

    def test_workbook_bytes_budget_marks_partial(self):
        # 全局"实际解压读取字节"预算：file_size 可伪造，实际计数才算数；
        # 触顶留痕并降为 partial（预算调小以便测试）
        import unittest.mock
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "大表"
        for r in range(1, 200):
            ws.append([f"值{r}"] * 8)
        buf = io.BytesIO()
        wb.save(buf)
        with unittest.mock.patch.object(xlsx_parser,
                                        "_MAX_MERGE_SCAN_TOTAL_BYTES", 256):
            merges, notes, incomplete = xlsx_parser._merged_ranges_by_sheet(
                buf.getvalue())
        self.assertFalse(merges)
        self.assertTrue(incomplete)
        self.assertTrue(any("字节预算" in n for n in notes),
                        f"字节预算触顶未留痕：{notes}")
        with unittest.mock.patch.object(xlsx_parser,
                                        "_MAX_MERGE_SCAN_TOTAL_BYTES", 256):
            rec = xlsx_parser.parse(buf.getvalue())
        self.assertEqual(rec["status"], "partial")
        self.assertTrue(rec["counts"]["merged_ranges_incomplete"])

    def test_missing_sheet_relationship_marks_incomplete(self):
        # workbook.xml 引用的关系缺失：不得静默跳过，置 incomplete 并留痕
        import zipfile as _zipfile
        good = _wb_with_merges(["A2:A3"])
        src = {}
        with _zipfile.ZipFile(io.BytesIO(good)) as z:
            for name in z.namelist():
                src[name] = z.read(name)
        src["xl/_rels/workbook.xml.rels"] = (
            '<?xml version="1.0"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/'
            'package/2006/relationships"/>').encode("utf-8")
        buf = io.BytesIO()
        with _zipfile.ZipFile(buf, "w") as z:
            for name, payload in src.items():
                z.writestr(name, payload)
        merges, notes, incomplete = xlsx_parser._merged_ranges_by_sheet(
            buf.getvalue())
        self.assertFalse(merges)
        self.assertTrue(incomplete)
        self.assertTrue(any("关系" in n or "缺失" in n for n in notes),
                        f"关系缺失未留痕：{notes}")


if __name__ == "__main__":
    unittest.main()
