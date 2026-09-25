# -*- coding: utf-8 -*-
"""标镜候选层合并单元格锚点回填回归（C 线审查 P2-8 下半）。

口径（主任务实现，承接解析层 xlsx_merged_range 披露）：
  - 纵向合并的投标人列：数据行主体为空但落在合并区域内时，以锚点值回填
    归属，证据回指锚点单元格证据 ID，组 note 披露"来自合并区域锚点"；
    缺归属不再静默看似完整；
  - 横向合并的表头：被覆盖列按锚点标签识别字段；若与首列形成同字段多列，
    走既有重复字段保守策略（unit_price 不自动取值并披露），绝不猜列；
  - 锚点值也为空（如公式无缓存）→ 维持不归属，不猜；
  - 无合并区域（旧解析输出）时行为与既有完全一致。

运行：
    cd app && python3 -W error::ResourceWarning -m unittest \\
        biaojing.tests.test_candidate_merge_fill -v
"""

from __future__ import annotations

import io
import unittest

import openpyxl

from biaojing import candidates, xlsx_parser


def _pipeline(wb) -> tuple[list[dict], list[dict]]:
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    rec = xlsx_parser.parse(data)
    tagged = candidates.assign_evidence_ids("sha", rec["evidence"])
    return tagged, candidates.extract_candidates("sha", "xlsx", tagged)


def _price_line_candidates(cands: list[dict]) -> list[dict]:
    return [c for c in cands if c["field"] == "price_lines"]


class CandidateMergeFillTests(unittest.TestCase):

    def test_vertical_bidder_merge_fills_ownership(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "报价"
        ws.append(["投标人", "综合单价(元)", "清单名称", "规格型号"])
        ws["A2"] = "甲公司"
        ws.merge_cells("A2:A4")
        for row, (price, item, spec) in zip(
                (2, 3, 4),
                ((100, "电线", "BV"), (110, "开关", "10A"),
                 (120, "水泵", "Q=10"))):
            ws.cell(row=row, column=2, value=price)
            ws.cell(row=row, column=3, value=item)
            ws.cell(row=row, column=4, value=spec)
        tagged, cands = _pipeline(wb)
        groups = _price_line_candidates(cands)
        self.assertEqual(len(groups), 1)
        value = groups[0]["value"]
        self.assertEqual(len(value), 3)
        self.assertEqual([line["unit_price"] for line in value],
                         ["100", "110", "120"])
        anchor = [e for e in tagged
                  if e.get("kind") == "xlsx_cell" and e.get("cell") == "A2"]
        self.assertTrue(anchor)
        self.assertIn(anchor[0]["evidence_id"], groups[0]["evidence_ids"])
        note = str(groups[0].get("note") or "")
        self.assertIn("合并", note)

    def test_horizontal_merged_header_is_conservative(self):
        # 合并表头 B1:C1 同为综合单价(元)：两列都识别为该字段，
        # 但绝不自动选列——unit_price=None 并披露，须人工指定
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "合并表头"
        ws["A1"] = "投标人"
        ws["B1"] = "综合单价(元)"
        ws.merge_cells("B1:C1")
        ws["D1"] = "清单名称"
        ws["E1"] = "规格型号"
        ws.append(["甲公司", 100, 888, "电线", "BV"])
        tagged, cands = _pipeline(wb)
        groups = _price_line_candidates(cands)
        self.assertEqual(len(groups), 1)
        lines = groups[0]["value"]
        self.assertEqual(len(lines), 1)
        self.assertIsNone(lines[0]["unit_price"])
        self.assertEqual(lines[0]["item_name"], "电线")
        note = str(groups[0].get("note") or "")
        self.assertIn("B", note)
        self.assertIn("C", note)

    def test_empty_anchor_keeps_row_unowned(self):
        # 锚点为公式且无缓存：维持不归属，不猜
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "空锚点"
        ws.append(["投标人", "综合单价(元)", "清单名称", "规格型号"])
        ws["A2"] = "=Z9"
        ws.merge_cells("A2:A3")
        ws["B2"] = 100
        ws["C2"] = "电线"
        ws["D2"] = "BV"
        ws["B3"] = 110
        ws["C3"] = "开关"
        ws["D3"] = "10A"
        _, cands = _pipeline(wb)
        self.assertEqual(_price_line_candidates(cands), [])

    def test_no_merge_behavior_unchanged(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "报价"
        ws.append(["投标人", "综合单价(元)", "清单名称", "规格型号"])
        ws.append(["甲公司", 100, "电线", "BV"])
        ws.append(["乙公司", 200, "开关", "10A"])
        _, cands = _pipeline(wb)
        groups = _price_line_candidates(cands)
        self.assertEqual(len(groups), 2)
        prices = sorted(line["unit_price"] for g in groups
                        for line in g["value"])
        self.assertEqual(prices, ["100", "200"])


if __name__ == "__main__":
    unittest.main()
