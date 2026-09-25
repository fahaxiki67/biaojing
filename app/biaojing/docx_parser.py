# -*- coding: utf-8 -*-
"""DOCX 段落、表格与本地内嵌图片 OCR，保留结构定位。"""

from __future__ import annotations

import io
import zipfile

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from . import pdf_parser

PARSER_NAME = f"python-docx {docx.__version__}"
EXTRACT_METHOD = (
    "python-docx 按文档体 XML 顺序抽取段落、表格、页眉页脚；"
    "本机 Tesseract OCR 识别内嵌 PNG/JPEG/GIF/BMP，保留图片所在结构位置"
)
OCR_MAX_TOTAL_IMAGE_BYTES = 100_000_000
EXTRACTION_VERSION = (
    f"{PARSER_NAME}; embedded image OCR; max={pdf_parser.OCR_MAX_IMAGE_BYTES}B/"
    f"{pdf_parser.OCR_MAX_PIXELS}px; total={OCR_MAX_TOTAL_IMAGE_BYTES}B; "
    f"timeout={pdf_parser.OCR_PAGE_TIMEOUT}s; v1"
)
_IMAGE_TAGS = {qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"}


def parse(data: bytes, progress=None, cancelled=None) -> dict:
    """返回文本/图片证据；所有 OCR 均为机器识别，须对照原件复核。"""
    document = docx.Document(io.BytesIO(data))
    notes = []
    evidence = []
    placements = []
    para_total = para_nonempty = table_total = cell_total = hf_total = 0
    para_no = table_no = image_no = 0

    def add_images(container, part, locator, **coords):
        nonlocal image_no
        local_no = 0
        for node in container.iter():
            if node.tag not in _IMAGE_TAGS:
                continue
            local_no += 1
            image_no += 1
            display = f"{locator} image {local_no}"
            item = {
                "kind": "docx_image", "locator": display,
                "image_index": image_no, "ocr_used": False,
                "page_status": "pending_ocr",
                "page_error": "图片尚未 OCR", **coords,
            }
            evidence.append(item)
            rid = (node.get(qn("r:embed")) or node.get(qn("r:id"))
                   or node.get(qn("r:link")))
            placements.append((item, part, rid,
                               bool(node.get(qn("r:link")) and not node.get(qn("r:embed")))))

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para_no += 1
            para_total += 1
            para = Paragraph(child, document)
            text = para.text
            if text and text.strip():
                para_nonempty += 1
                evidence.append({"kind": "docx_paragraph",
                                 "locator": f"paragraph {para_no}",
                                 "paragraph_index": para_no, "text": text})
            add_images(child, para.part, f"paragraph {para_no}",
                       paragraph_index=para_no)
        elif child.tag == qn("w:tbl"):
            table_no += 1
            table_total += 1
            table = Table(child, document)
            for r, row in enumerate(table.rows, start=1):
                for c, cell in enumerate(row.cells, start=1):
                    cell_total += 1
                    text = cell.text
                    if text and text.strip():
                        evidence.append({
                            "kind": "docx_table_cell",
                            "locator": f"table {table_no}!R{r}C{c}",
                            "table_index": table_no, "row": r, "col": c,
                            "text": text,
                        })
                    for para in cell.paragraphs:
                        add_images(para._p, para.part,
                                   f"table {table_no}!R{r}C{c}",
                                   table_index=table_no, row=r, col=c)

    for s_no, section in enumerate(document.sections, start=1):
        for kind, part in (("header", section.header), ("footer", section.footer)):
            try:
                hf_paras = list(part.paragraphs)
                hf_tables = list(part.tables)
            except Exception:
                notes.append(f"第 {s_no} 节 {kind} 读取失败，跳过")
                continue
            for p_no, para in enumerate(hf_paras, start=1):
                hf_total += 1
                text = para.text
                if text and text.strip():
                    evidence.append({
                        "kind": f"docx_{kind}_paragraph",
                        "locator": f"section {s_no} {kind} paragraph {p_no}",
                        "section_index": s_no, "paragraph_index": p_no,
                        "text": text,
                    })
                add_images(para._p, para.part,
                           f"section {s_no} {kind} paragraph {p_no}",
                           section_index=s_no, paragraph_index=p_no)
            for t_no, table in enumerate(hf_tables, start=1):
                for r, row in enumerate(table.rows, start=1):
                    for c, cell in enumerate(row.cells, start=1):
                        text = cell.text
                        if text and text.strip():
                            evidence.append({
                                "kind": f"docx_{kind}_table_cell",
                                "locator": f"section {s_no} {kind} table {t_no}!R{r}C{c}",
                                "section_index": s_no, "table_index": t_no,
                                "row": r, "col": c, "text": text,
                            })
                        for para in cell.paragraphs:
                            add_images(
                                para._p, para.part,
                                f"section {s_no} {kind} table {t_no}!R{r}C{c}",
                                section_index=s_no, table_index=t_no,
                                row=r, col=c)

    image_total = len(placements)
    images_ocr = images_pending = 0
    cancelled_run = False
    if placements:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            part_sizes = {info.filename: info.file_size
                          for info in archive.infolist() if not info.is_dir()}
        binary, lang, unavailable = pdf_parser._ocr_runtime()
        if unavailable:
            notes.append(unavailable.replace("扫描页", "Word 图片"))
        cache = {}
        decoded_bytes = 0
        for index, (item, owner, rid, external) in enumerate(placements, start=1):
            if cancelled and cancelled():
                cancelled_run = True
                for pending, _, _, _ in placements[index - 1:]:
                    pending["page_error"] = "用户取消处理；此图片尚未 OCR，可重新导入"
                    images_pending += 1
                notes.append(
                    f"用户取消处理；已识别 {index - 1}/{image_total} 张图片，"
                    "其余图片保留待 OCR")
                break
            if progress:
                progress(index, image_total, "processing", None)
            key = (id(owner), rid)
            if key in cache:
                text, error = cache[key]
            elif external:
                text, error = "", "外部链接图片未下载；未访问网络"
                cache[key] = (text, error)
            elif not rid:
                text, error = "", "图片缺少 OOXML 关系标识"
                cache[key] = (text, error)
            else:
                try:
                    image_part = owner.related_parts[rid]
                except Exception:
                    image_part = None
                if image_part is None or not image_part.content_type.startswith("image/"):
                    text, error = "", "图片关系无法解析"
                    cache[key] = (text, error)
                else:
                    part_name = str(image_part.partname).lstrip("/")
                    size = part_sizes.get(part_name)
                    if size is None:
                        text, error = "", "图片数据未在 DOCX 容器中找到"
                    elif size > pdf_parser.OCR_MAX_IMAGE_BYTES:
                        text, error = "", "图片字节数超出 OCR 上限"
                    elif decoded_bytes + size > OCR_MAX_TOTAL_IMAGE_BYTES:
                        text, error = "", "本文件图片 OCR 数据总量超出上限"
                    elif unavailable:
                        text, error = "", unavailable.replace("扫描页", "Word 图片")
                    else:
                        try:
                            image_data = image_part.blob
                            decoded_bytes += size
                            if len(image_data) != size:
                                text, error = "", "图片解压字节数与 DOCX 声明不符"
                            else:
                                text, error = pdf_parser.ocr_image(
                                    image_data, binary, lang)
                        except Exception as exc:
                            text, error = "", f"图片读取失败（{type(exc).__name__}）"
                    cache[key] = (text, error)
            item["text"] = text
            if text:
                item.update(ocr_used=True, page_status="ocr",
                            page_error=None,
                            page_note="机器 OCR 文本，请对照原始 Word 图片复核")
                images_ocr += 1
            else:
                item.update(page_status="pending_ocr",
                            page_error=error or "图片仍待 OCR")
                images_pending += 1
            if progress:
                progress(index, image_total, "page_done", item["page_status"])
        if images_ocr:
            notes.append("Word 内嵌图片文字由本机 OCR 机器识别，须对照原件人工复核")

    counts = {
        "paragraphs_total": para_total,
        "paragraphs_nonempty": para_nonempty,
        "tables_total": table_total,
        "table_cells_total": cell_total,
        "header_footer_paragraphs_total": hf_total,
        "inline_images": image_total,
        "images_ocr": images_ocr,
        "images_pending_ocr": images_pending,
    }
    core = document.core_properties
    metadata = {k: v for k, v in {
        "title": core.title or "", "author": core.author or "",
        "last_modified_by": core.last_modified_by or "",
        "created": core.created.isoformat() if core.created else "",
        "modified": core.modified.isoformat() if core.modified else "",
        "revision": core.revision if core.revision else 0,
    }.items() if v not in ("", 0)}

    readable = any(e.get("kind") != "docx_image" and
                   str(e.get("text") or "").strip() for e in evidence)
    readable = readable or images_ocr > 0
    if not readable and images_pending:
        status = "pending_ocr"
    elif readable and images_pending:
        status = "partial"
    elif readable:
        status = "success"
    else:
        return {
            "ok": False, "status": "failed",
            "error": "DOCX 无任何文本内容（正文、表格、页眉页脚均为空）",
            "evidence": evidence, "counts": counts,
            "metadata": metadata, "notes": notes,
            "cancelled": cancelled_run,
        }
    return {
        "ok": True, "status": status, "error": None,
        "evidence": evidence, "counts": counts, "metadata": metadata,
        "notes": notes, "cancelled": cancelled_run,
    }
